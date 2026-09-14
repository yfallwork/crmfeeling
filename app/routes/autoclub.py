import csv
import io
import json
import os
import re
import uuid
from datetime import date, datetime, timedelta
import requests
from flask import (Blueprint, render_template, request, redirect,
                   url_for, flash, Response, current_app, jsonify, abort)
from flask_login import login_required, current_user
from werkzeug.utils import secure_filename
from app.extensions import db
from app.models.socio import Socio
from app.models.patrocinador import Patrocinador
from app.models.autoclub_nota import AutoclubNota
from app.models.piloto import Piloto
from app.models.vehiculo import Vehiculo
from app.models.inscripcion import Inscripcion
from app.models.inscripcion_piloto import InscripcionPiloto
from app.models.gasto_inscripcion import GastoInscripcion, CATEGORIAS_GASTO
from app.models.staff_miembro import StaffMiembro, TIPOS_STAFF
from app.models.resultado_inscripcion import ResultadoInscripcion
from app.models.preinscripcion_carcross import PreinscripcionCarcross, ESTADOS_PREINSCRIPCION, EXPERIENCIA_PREVIA_OPCIONES
from app.models.firmable import SesionFirmables, DocumentoFirmado
from app.models.entrevista import EntrevistaPiloto, BLOQUE_1_FAMILIA, BLOQUE_2_CANDIDATA, CLAVES_PREGUNTAS
from app.services.firmables_textos import TIPOS_FIRMABLE, TIPOS_FIRMABLE_KEYS, TIPOS_FIRMABLE_LABELS
from app.services.log_service import registrar_log

autoclub_bp = Blueprint("autoclub", __name__)

TIPOS_SOCIO = ["regular", "premium", "familiar", "vip"]
LOGO_EXTENSIONS = {"png", "jpg", "jpeg", "gif", "webp", "svg"}
DOC_EXTENSIONS = {"pdf", "jpg", "jpeg", "png", "webp", "doc", "docx"}
FOTO_PRENSA_EXTENSIONS = {"png", "jpg", "jpeg", "gif", "webp"}
FOTO_PRENSA_MIME = {
    "png": "png", "jpg": "jpeg", "jpeg": "jpeg", "gif": "gif", "webp": "webp",
}
ESTADOS_INSCRIPCION = ["pendiente", "enviada", "aceptada", "rechazada"]

SECTORES_ESPAÑA = sorted([
    "Administración pública",
    "Aeronáutica y aviación",
    "Agricultura y ganadería",
    "Alimentación y bebidas",
    "Arquitectura y urbanismo",
    "Arte y cultura",
    "Asesoría y gestoría",
    "Automoción y motor",
    "Banca y finanzas",
    "Biotecnología",
    "Comercio al por menor",
    "Comercio al por mayor",
    "Comunicación y marketing",
    "Construcción e inmobiliaria",
    "Consultoría y estrategia",
    "Defensa y seguridad pública",
    "Deporte y ocio activo",
    "Diseño gráfico y UX",
    "Distribución y logística",
    "Educación y formación",
    "Electrónica y semiconductores",
    "Energía y utilities",
    "Entretenimiento y espectáculos",
    "Envases y embalajes",
    "Eventos y producción",
    "Farmacia y salud",
    "Fotografía y audiovisual",
    "Gastronomía y restauración",
    "Hostelería y turismo",
    "Industria química",
    "Industria manufacturera",
    "Ingeniería e I+D",
    "Joyería y moda de lujo",
    "Legal y compliance",
    "Maquinaria e industria pesada",
    "Medioambiente y sostenibilidad",
    "Medios de comunicación",
    "Minería y extracción",
    "Moda y textil",
    "Naval y marítimo",
    "Nuevas tecnologías e IA",
    "Óptica y oftalmología",
    "Papel e impresión",
    "Publicidad y RRPP",
    "Recursos humanos",
    "Salud y bienestar",
    "Seguros",
    "Seguridad privada",
    "Servicios sociales",
    "Software y SaaS",
    "Telecomunicaciones",
    "Transporte y movilidad",
    "Veterinaria y mascotas",
    "Viajes y agencias de viaje",
    "Videojuegos y esports",
])


# ── CONTEXT PROCESSOR — inyecta sectores en todas las plantillas del blueprint ─

@autoclub_bp.context_processor
def inject_autoclub():
    return {"sectores_espana": SECTORES_ESPAÑA}


# ── HELPERS UPLOAD ─────────────────────────────────────────────────────────────

def _logo_dir():
    path = os.path.join(current_app.root_path, "static", "uploads", "autoclub", "patrocinadores")
    os.makedirs(path, exist_ok=True)
    return path


def _save_logo(file, old_filename=None):
    if not file or not file.filename:
        return None
    ext = file.filename.rsplit(".", 1)[-1].lower()
    if ext not in LOGO_EXTENSIONS:
        return None
    filename = f"{uuid.uuid4().hex}.{ext}"
    file.save(os.path.join(_logo_dir(), filename))
    if old_filename:
        _delete_logo(old_filename)
    return filename


def _delete_logo(filename):
    if not filename:
        return
    try:
        os.remove(os.path.join(_logo_dir(), filename))
    except OSError:
        pass


# ── DASHBOARD ──────────────────────────────────────────────────────────────────

@autoclub_bp.route("/")
@login_required
def dashboard():
    total      = Socio.query.count()
    activos    = Socio.query.filter_by(activo=True).count()
    inactivos  = Socio.query.filter_by(activo=False).count()
    por_tipo   = {t: Socio.query.filter_by(tipo=t, activo=True).count() for t in TIPOS_SOCIO}
    recientes  = Socio.query.filter_by(activo=True).order_by(Socio.creado_en.desc()).limit(6).all()

    total_patrocinadores   = Patrocinador.query.count()
    activos_patrocinadores = Patrocinador.query.filter_by(activo=True).count()
    recientes_patrocinadores = Patrocinador.query.filter_by(activo=True).order_by(Patrocinador.creado_en.desc()).limit(4).all()

    return render_template(
        "autoclub/dashboard.html",
        total=total, activos=activos, inactivos=inactivos,
        por_tipo=por_tipo, recientes=recientes, tipos=TIPOS_SOCIO,
        total_patrocinadores=total_patrocinadores,
        activos_patrocinadores=activos_patrocinadores,
        recientes_patrocinadores=recientes_patrocinadores,
    )


# ── BÚSQUEDA (AJAX) ────────────────────────────────────────────────────────────

@autoclub_bp.route("/buscar")
@login_required
def buscar():
    q = request.args.get("q", "").strip()
    if len(q) < 2:
        return jsonify({"socios": [], "patrocinadores": []})

    like = f"%{q}%"

    socios = Socio.query.filter(db.or_(
        Socio.nombre.ilike(like),
        Socio.apellido.ilike(like),
        Socio.email.ilike(like),
        Socio.dni.ilike(like),
    )).limit(6).all()

    patrocinadores = Patrocinador.query.filter(db.or_(
        Patrocinador.nombre.ilike(like),
        Patrocinador.sector.ilike(like),
        Patrocinador.persona_contacto.ilike(like),
        Patrocinador.email.ilike(like),
    )).limit(6).all()

    return jsonify({
        "socios": [
            {
                "id": s.id,
                "nombre": s.nombre_completo,
                "tipo": s.tipo,
                "email": s.email,
                "activo": s.activo,
                "url": url_for("autoclub.detalle", id=s.id),
            }
            for s in socios
        ],
        "patrocinadores": [
            {
                "id": p.id,
                "nombre": p.nombre,
                "sector": p.sector,
                "tiene_logo": p.tiene_logo,
                "logo": p.logo_principal,
                "url": url_for("autoclub.patrocinadores_detalle", id=p.id),
            }
            for p in patrocinadores
        ],
    })


# ══════════════════════════════════════════════════════════════════════════════
#  SOCIOS
# ══════════════════════════════════════════════════════════════════════════════

@autoclub_bp.route("/socios/")
@login_required
def lista():
    q      = request.args.get("q", "").strip()
    estado = request.args.get("estado", "activos")
    tipo   = request.args.get("tipo", "")
    page   = request.args.get("page", 1, type=int)

    query = Socio.query
    if estado == "activos":
        query = query.filter_by(activo=True)
    elif estado == "inactivos":
        query = query.filter_by(activo=False)
    if q:
        like = f"%{q}%"
        query = query.filter(db.or_(
            Socio.nombre.ilike(like), Socio.apellido.ilike(like),
            Socio.email.ilike(like), Socio.telefono.ilike(like),
            Socio.dni.ilike(like),
        ))
    if tipo:
        query = query.filter_by(tipo=tipo)

    socios = query.order_by(Socio.nombre.asc()).paginate(page=page, per_page=25, error_out=False)
    return render_template("autoclub/socios/lista.html",
                           socios=socios, q=q, estado=estado, tipo=tipo, tipos=TIPOS_SOCIO)


@autoclub_bp.route("/socios/nuevo", methods=["GET", "POST"])
@login_required
def nuevo():
    if request.method == "POST":
        fecha_alta = None
        raw = request.form.get("fecha_alta", "").strip()
        if raw:
            try:
                fecha_alta = date.fromisoformat(raw)
            except ValueError:
                pass

        socio = Socio(
            nombre=request.form.get("nombre", "").strip(),
            apellido=request.form.get("apellido", "").strip(),
            email=request.form.get("email", "").strip().lower(),
            telefono=request.form.get("telefono", "").strip(),
            dni=request.form.get("dni", "").strip(),
            tipo=request.form.get("tipo", "regular"),
            fecha_alta=fecha_alta,
            notas=request.form.get("notas", "").strip(),
        )
        db.session.add(socio)
        db.session.commit()
        registrar_log("crear", "socio", socio.id, f"Socio creado: {socio.nombre_completo}")
        flash(f"Socio {socio.nombre_completo} creado correctamente.", "success")
        return redirect(url_for("autoclub.detalle", id=socio.id))

    return render_template("autoclub/socios/form.html", socio=None, datos={}, tipos=TIPOS_SOCIO)


@autoclub_bp.route("/socios/<int:id>")
@login_required
def detalle(id):
    socio = Socio.query.get_or_404(id)
    notas = AutoclubNota.query.filter_by(entidad="socio", entidad_id=id).order_by(AutoclubNota.creado_en.desc()).all()
    return render_template("autoclub/socios/detalle.html", socio=socio, notas=notas)


@autoclub_bp.route("/socios/<int:id>/editar", methods=["GET", "POST"])
@login_required
def editar(id):
    socio = Socio.query.get_or_404(id)
    if request.method == "POST":
        fecha_alta = socio.fecha_alta
        raw = request.form.get("fecha_alta", "").strip()
        if raw:
            try:
                fecha_alta = date.fromisoformat(raw)
            except ValueError:
                pass
        socio.nombre     = request.form.get("nombre", "").strip()
        socio.apellido   = request.form.get("apellido", "").strip()
        socio.email      = request.form.get("email", "").strip().lower()
        socio.telefono   = request.form.get("telefono", "").strip()
        socio.dni        = request.form.get("dni", "").strip()
        socio.tipo       = request.form.get("tipo", "regular")
        socio.fecha_alta = fecha_alta
        socio.notas      = request.form.get("notas", "").strip()
        db.session.commit()
        registrar_log("editar", "socio", socio.id, f"Socio editado: {socio.nombre_completo}")
        flash("Socio actualizado correctamente.", "success")
        return redirect(url_for("autoclub.detalle", id=socio.id))
    return render_template("autoclub/socios/form.html", socio=socio, datos={}, tipos=TIPOS_SOCIO)


@autoclub_bp.route("/socios/<int:id>/estado", methods=["POST"])
@login_required
def cambiar_estado(id):
    socio = Socio.query.get_or_404(id)
    if socio.activo:
        socio.desactivar()
        registrar_log("cambiar_estado", "socio", socio.id, f"Socio desactivado: {socio.nombre_completo}")
        flash(f"{socio.nombre_completo} marcado como inactivo.", "warning")
    else:
        socio.activar()
        registrar_log("cambiar_estado", "socio", socio.id, f"Socio reactivado: {socio.nombre_completo}")
        flash(f"{socio.nombre_completo} reactivado correctamente.", "success")
    db.session.commit()
    return redirect(url_for("autoclub.detalle", id=socio.id))


@autoclub_bp.route("/socios/<int:id>/eliminar", methods=["POST"])
@login_required
def eliminar(id):
    socio = Socio.query.get_or_404(id)
    nombre = socio.nombre_completo
    db.session.delete(socio)
    db.session.commit()
    registrar_log("eliminar", "socio", id, f"Socio eliminado: {nombre}")
    flash(f"Socio {nombre} eliminado permanentemente.", "info")
    return redirect(url_for("autoclub.lista"))


@autoclub_bp.route("/socios/exportar")
@login_required
def exportar():
    q = request.args.get("q", "").strip()
    estado = request.args.get("estado", "")
    tipo   = request.args.get("tipo", "")
    query  = Socio.query
    if estado == "activos":
        query = query.filter_by(activo=True)
    elif estado == "inactivos":
        query = query.filter_by(activo=False)
    if q:
        like = f"%{q}%"
        query = query.filter(db.or_(
            Socio.nombre.ilike(like), Socio.apellido.ilike(like),
            Socio.email.ilike(like), Socio.telefono.ilike(like),
            Socio.dni.ilike(like),
        ))
    if tipo:
        query = query.filter_by(tipo=tipo)
    socios = query.order_by(Socio.nombre.asc()).all()
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(["ID", "Nombre", "Apellido", "Email", "Teléfono",
                     "DNI", "Tipo", "Fecha_alta", "Activo", "Registrado"])
    for s in socios:
        writer.writerow([
            s.id, s.nombre, s.apellido, s.email, s.telefono, s.dni, s.tipo,
            s.fecha_alta.strftime("%d/%m/%Y") if s.fecha_alta else "",
            "Sí" if s.activo else "No",
            s.creado_en.strftime("%d/%m/%Y") if s.creado_en else "",
        ])
    return Response(
        "﻿" + output.getvalue(),
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=socios_autoclub.csv"},
    )


# ══════════════════════════════════════════════════════════════════════════════
#  PATROCINADORES
# ══════════════════════════════════════════════════════════════════════════════

def _update_logo_field(p, field_attr, file_field):
    """Guarda el logo subido para un campo específico. Retorna el nuevo filename o None."""
    file = request.files.get(file_field)
    new_fn = _save_logo(file, old_filename=getattr(p, field_attr) if p else None)
    if new_fn:
        return new_fn
    # Opción de borrar sin reemplazar
    if p and request.form.get(f"eliminar_{file_field}") == "1":
        _delete_logo(getattr(p, field_attr))
        return ""
    return None  # sin cambio


def _extract_logos_from_form(p=None):
    """Procesa los 3 logos del formulario. Devuelve dict con los campos que cambian."""
    changes = {}
    for attr, field in [
        ("logo_positivo_filename", "logo_positivo"),
        ("logo_negativo_filename", "logo_negativo"),
        ("logo_banner_filename",   "logo_banner"),
    ]:
        result = _update_logo_field(p, attr, field)
        if result is not None:
            changes[attr] = result
    return changes


@autoclub_bp.route("/patrocinadores/")
@login_required
def patrocinadores_lista():
    q      = request.args.get("q", "").strip()
    estado = request.args.get("estado", "activos")
    page   = request.args.get("page", 1, type=int)

    query = Patrocinador.query
    if estado == "activos":
        query = query.filter_by(activo=True)
    elif estado == "inactivos":
        query = query.filter_by(activo=False)
    if q:
        like = f"%{q}%"
        query = query.filter(db.or_(
            Patrocinador.nombre.ilike(like),
            Patrocinador.sector.ilike(like),
            Patrocinador.persona_contacto.ilike(like),
            Patrocinador.email.ilike(like),
        ))

    patrocinadores = query.order_by(Patrocinador.nombre.asc()).paginate(
        page=page, per_page=25, error_out=False
    )
    return render_template("autoclub/patrocinadores/lista.html",
                           patrocinadores=patrocinadores, q=q, estado=estado)


@autoclub_bp.route("/patrocinadores/nuevo", methods=["GET", "POST"])
@login_required
def patrocinadores_nuevo():
    if request.method == "POST":
        logos = _extract_logos_from_form()
        p = Patrocinador(
            nombre           = request.form.get("nombre", "").strip(),
            sector           = request.form.get("sector", "").strip(),
            persona_contacto = request.form.get("persona_contacto", "").strip(),
            email            = request.form.get("email", "").strip().lower(),
            telefono         = request.form.get("telefono", "").strip(),
            web              = request.form.get("web", "").strip(),
            instagram        = request.form.get("instagram", "").strip(),
            facebook         = request.form.get("facebook", "").strip(),
            twitter          = request.form.get("twitter", "").strip(),
            linkedin         = request.form.get("linkedin", "").strip(),
            tiktok           = request.form.get("tiktok", "").strip(),
            youtube          = request.form.get("youtube", "").strip(),
            logo_positivo_filename = logos.get("logo_positivo_filename", ""),
            logo_negativo_filename = logos.get("logo_negativo_filename", ""),
            logo_banner_filename   = logos.get("logo_banner_filename", ""),
            notas = request.form.get("notas", "").strip(),
        )
        db.session.add(p)
        db.session.commit()
        registrar_log("crear", "patrocinador", p.id, f"Patrocinador creado: {p.nombre}")
        flash(f"Patrocinador {p.nombre} creado correctamente.", "success")
        return redirect(url_for("autoclub.patrocinadores_detalle", id=p.id))

    return render_template("autoclub/patrocinadores/form.html", p=None, datos={})


@autoclub_bp.route("/patrocinadores/<int:id>")
@login_required
def patrocinadores_detalle(id):
    p = Patrocinador.query.get_or_404(id)
    notas = AutoclubNota.query.filter_by(entidad="patrocinador", entidad_id=id).order_by(AutoclubNota.creado_en.desc()).all()
    return render_template("autoclub/patrocinadores/detalle.html", p=p, notas=notas)


@autoclub_bp.route("/patrocinadores/<int:id>/editar", methods=["GET", "POST"])
@login_required
def patrocinadores_editar(id):
    p = Patrocinador.query.get_or_404(id)
    if request.method == "POST":
        logos = _extract_logos_from_form(p)
        for attr, val in logos.items():
            setattr(p, attr, val)

        p.nombre           = request.form.get("nombre", "").strip()
        p.sector           = request.form.get("sector", "").strip()
        p.persona_contacto = request.form.get("persona_contacto", "").strip()
        p.email            = request.form.get("email", "").strip().lower()
        p.telefono         = request.form.get("telefono", "").strip()
        p.web              = request.form.get("web", "").strip()
        p.instagram        = request.form.get("instagram", "").strip()
        p.facebook         = request.form.get("facebook", "").strip()
        p.twitter          = request.form.get("twitter", "").strip()
        p.linkedin         = request.form.get("linkedin", "").strip()
        p.tiktok           = request.form.get("tiktok", "").strip()
        p.youtube          = request.form.get("youtube", "").strip()
        p.notas            = request.form.get("notas", "").strip()
        db.session.commit()
        registrar_log("editar", "patrocinador", p.id, f"Patrocinador editado: {p.nombre}")
        flash("Patrocinador actualizado correctamente.", "success")
        return redirect(url_for("autoclub.patrocinadores_detalle", id=p.id))

    return render_template("autoclub/patrocinadores/form.html", p=p, datos={})


@autoclub_bp.route("/patrocinadores/<int:id>/estado", methods=["POST"])
@login_required
def patrocinadores_estado(id):
    p = Patrocinador.query.get_or_404(id)
    if p.activo:
        p.activo = False
        registrar_log("cambiar_estado", "patrocinador", p.id, f"Patrocinador desactivado: {p.nombre}")
        flash(f"{p.nombre} marcado como inactivo.", "warning")
    else:
        p.activo = True
        registrar_log("cambiar_estado", "patrocinador", p.id, f"Patrocinador reactivado: {p.nombre}")
        flash(f"{p.nombre} reactivado correctamente.", "success")
    db.session.commit()
    return redirect(url_for("autoclub.patrocinadores_detalle", id=p.id))


@autoclub_bp.route("/patrocinadores/<int:id>/eliminar", methods=["POST"])
@login_required
def patrocinadores_eliminar(id):
    p = Patrocinador.query.get_or_404(id)
    nombre = p.nombre
    for fn in [p.logo_positivo_filename, p.logo_negativo_filename, p.logo_banner_filename]:
        _delete_logo(fn)
    db.session.delete(p)
    db.session.commit()
    registrar_log("eliminar", "patrocinador", id, f"Patrocinador eliminado: {nombre}")
    flash(f"Patrocinador {nombre} eliminado.", "info")
    return redirect(url_for("autoclub.patrocinadores_lista"))


# ══════════════════════════════════════════════════════════════════════════════
#  DOCS UPLOAD HELPER
# ══════════════════════════════════════════════════════════════════════════════

DOCS_FOLDER = os.path.join(os.path.dirname(__file__), '..', 'static', 'uploads', 'autoclub', 'docs')


def _save_doc(file):
    if not file or not file.filename:
        return None
    ext = os.path.splitext(secure_filename(file.filename))[1].lstrip(".").lower()
    if ext not in DOC_EXTENSIONS:
        return None
    os.makedirs(DOCS_FOLDER, exist_ok=True)
    fn = f"{uuid.uuid4().hex}.{ext}"
    file.save(os.path.join(DOCS_FOLDER, fn))
    return fn


def _delete_doc(filename):
    if not filename:
        return
    try:
        os.remove(os.path.join(DOCS_FOLDER, filename))
    except OSError:
        pass


def _leer_doc_gasto(campo, etiqueta):
    """Lee un fichero de formulario (comprobante/factura) para un gasto.
    Si el usuario adjuntó algo con un formato no admitido, avisa con un
    flash en vez de descartarlo en silencio."""
    f = request.files.get(campo)
    if not f or not f.filename:
        return None
    fn = _save_doc(f)
    if not fn:
        flash(f"El/la {etiqueta} «{f.filename}» no se pudo adjuntar: formato no admitido "
              f"(usa PDF, JPG, PNG, WEBP, DOC o DOCX). El gasto se ha guardado sin ese archivo.", "warning")
    return fn


def _parse_date(raw):
    if not raw:
        return None
    try:
        from datetime import date as _date
        return _date.fromisoformat(raw.strip())
    except ValueError:
        return None


def _parse_int(raw):
    try:
        return int(raw)
    except (TypeError, ValueError):
        return None


def _parse_float(raw):
    try:
        return float(raw)
    except (TypeError, ValueError):
        return None


# ══════════════════════════════════════════════════════════════════════════════
#  PILOTOS
# ══════════════════════════════════════════════════════════════════════════════

@autoclub_bp.route("/pilotos")
@login_required
def pilotos_lista():
    q      = request.args.get("q", "").strip()
    estado = request.args.get("estado", "activos")
    page   = request.args.get("page", 1, type=int)

    query = Piloto.query
    if estado == "activos":
        query = query.filter_by(activo=True)
    elif estado == "inactivos":
        query = query.filter_by(activo=False)
    if q:
        like = f"%{q}%"
        query = query.filter(db.or_(
            Piloto.nombre.ilike(like),
            Piloto.apellidos.ilike(like),
            Piloto.dni.ilike(like),
            Piloto.email.ilike(like),
            Piloto.num_licencia.ilike(like),
        ))

    pilotos = query.order_by(Piloto.nombre.asc()).paginate(page=page, per_page=25, error_out=False)
    return render_template("autoclub/pilotos/lista.html",
                           pilotos=pilotos, q=q, estado=estado)


@autoclub_bp.route("/pilotos/nuevo", methods=["GET", "POST"])
@login_required
def pilotos_nuevo():
    if request.method == "POST":
        p = Piloto(
            nombre=request.form.get("nombre", "").strip(),
            apellidos=request.form.get("apellidos", "").strip(),
            dni=request.form.get("dni", "").strip(),
            fecha_nacimiento=_parse_date(request.form.get("fecha_nacimiento")),
            nacionalidad=request.form.get("nacionalidad", "").strip(),
            calle=request.form.get("calle", "").strip(),
            cp=request.form.get("cp", "").strip(),
            localidad=request.form.get("localidad", "").strip(),
            provincia=request.form.get("provincia", "").strip(),
            telefono=request.form.get("telefono", "").strip(),
            telefono_asistencia=request.form.get("telefono_asistencia", "").strip(),
            email=request.form.get("email", "").strip().lower(),
            num_licencia=request.form.get("num_licencia", "").strip(),
            tipo_licencia=request.form.get("tipo_licencia", "").strip(),
            federacion=request.form.get("federacion", "").strip(),
            num_licencia_concursante=request.form.get("num_licencia_concursante", "").strip(),
            nombre_concursante=request.form.get("nombre_concursante", "").strip(),
            grupo_sanguineo=request.form.get("grupo_sanguineo", "").strip(),
            alergias=request.form.get("alergias", "").strip(),
            contacto_emergencia_nombre=request.form.get("contacto_emergencia_nombre", "").strip(),
            contacto_emergencia_parentesco=request.form.get("contacto_emergencia_parentesco", "").strip(),
            contacto_emergencia_telefono=request.form.get("contacto_emergencia_telefono", "").strip(),
            tutor_nombre=request.form.get("tutor_nombre", "").strip(),
            tutor_dni=request.form.get("tutor_dni", "").strip(),
        )
        dni_doc = _save_doc(request.files.get("dni_doc"))
        if dni_doc:
            p.dni_doc_filename = dni_doc
        tutor_doc = _save_doc(request.files.get("tutor_doc"))
        if tutor_doc:
            p.tutor_doc_filename = tutor_doc
        db.session.add(p)
        db.session.commit()
        registrar_log("crear", "piloto", p.id, f"Piloto creado: {p.nombre_completo}")
        flash(f"Piloto {p.nombre_completo} creado correctamente.", "success")
        return redirect(url_for("autoclub.pilotos_detalle", id=p.id))
    return render_template("autoclub/pilotos/form.html", piloto=None)


@autoclub_bp.route("/pilotos/<int:id>")
@login_required
def pilotos_detalle(id):
    piloto = Piloto.query.get_or_404(id)
    return render_template("autoclub/pilotos/detalle.html", piloto=piloto)


@autoclub_bp.route("/pilotos/<int:id>/editar", methods=["GET", "POST"])
@login_required
def pilotos_editar(id):
    piloto = Piloto.query.get_or_404(id)
    if request.method == "POST":
        piloto.nombre = request.form.get("nombre", "").strip()
        piloto.apellidos = request.form.get("apellidos", "").strip()
        piloto.dni = request.form.get("dni", "").strip()
        piloto.fecha_nacimiento = _parse_date(request.form.get("fecha_nacimiento"))
        piloto.nacionalidad = request.form.get("nacionalidad", "").strip()
        piloto.calle = request.form.get("calle", "").strip()
        piloto.cp = request.form.get("cp", "").strip()
        piloto.localidad = request.form.get("localidad", "").strip()
        piloto.provincia = request.form.get("provincia", "").strip()
        piloto.telefono = request.form.get("telefono", "").strip()
        piloto.telefono_asistencia = request.form.get("telefono_asistencia", "").strip()
        piloto.email = request.form.get("email", "").strip().lower()
        piloto.num_licencia = request.form.get("num_licencia", "").strip()
        piloto.tipo_licencia = request.form.get("tipo_licencia", "").strip()
        piloto.federacion = request.form.get("federacion", "").strip()
        piloto.num_licencia_concursante = request.form.get("num_licencia_concursante", "").strip()
        piloto.nombre_concursante = request.form.get("nombre_concursante", "").strip()
        piloto.grupo_sanguineo = request.form.get("grupo_sanguineo", "").strip()
        piloto.alergias = request.form.get("alergias", "").strip()
        piloto.contacto_emergencia_nombre = request.form.get("contacto_emergencia_nombre", "").strip()
        piloto.contacto_emergencia_parentesco = request.form.get("contacto_emergencia_parentesco", "").strip()
        piloto.contacto_emergencia_telefono = request.form.get("contacto_emergencia_telefono", "").strip()
        piloto.tutor_nombre = request.form.get("tutor_nombre", "").strip()
        piloto.tutor_dni = request.form.get("tutor_dni", "").strip()
        dni_doc = _save_doc(request.files.get("dni_doc"))
        if dni_doc:
            _delete_doc(piloto.dni_doc_filename)
            piloto.dni_doc_filename = dni_doc
        tutor_doc = _save_doc(request.files.get("tutor_doc"))
        if tutor_doc:
            _delete_doc(piloto.tutor_doc_filename)
            piloto.tutor_doc_filename = tutor_doc
        db.session.commit()
        registrar_log("editar", "piloto", piloto.id, f"Piloto editado: {piloto.nombre_completo}")
        flash("Piloto actualizado correctamente.", "success")
        return redirect(url_for("autoclub.pilotos_detalle", id=piloto.id))
    return render_template("autoclub/pilotos/form.html", piloto=piloto)


@autoclub_bp.route("/pilotos/<int:id>/eliminar", methods=["POST"])
@login_required
def pilotos_eliminar(id):
    piloto = Piloto.query.get_or_404(id)
    nombre = piloto.nombre_completo
    _delete_doc(piloto.dni_doc_filename)
    _delete_doc(piloto.tutor_doc_filename)
    db.session.delete(piloto)
    db.session.commit()
    registrar_log("eliminar", "piloto", id, f"Piloto eliminado: {nombre}")
    flash(f"Piloto {nombre} eliminado.", "info")
    return redirect(url_for("autoclub.pilotos_lista"))


@autoclub_bp.route("/pilotos/<int:id>/toggle", methods=["POST"])
@login_required
def pilotos_toggle(id):
    piloto = Piloto.query.get_or_404(id)
    piloto.activo = not piloto.activo
    db.session.commit()
    estado = "reactivado" if piloto.activo else "desactivado"
    registrar_log("cambiar_estado", "piloto", piloto.id, f"Piloto {estado}: {piloto.nombre_completo}")
    flash(f"{piloto.nombre_completo} {estado} correctamente.", "success" if piloto.activo else "warning")
    return redirect(url_for("autoclub.pilotos_detalle", id=piloto.id))


# ══════════════════════════════════════════════════════════════════════════════
#  VEHÍCULOS
# ══════════════════════════════════════════════════════════════════════════════

@autoclub_bp.route("/vehiculos")
@login_required
def vehiculos_lista():
    q    = request.args.get("q", "").strip()
    tipo = request.args.get("tipo_vehiculo", "")
    page = request.args.get("page", 1, type=int)

    query = Vehiculo.query
    if tipo:
        query = query.filter_by(tipo_vehiculo=tipo)
    if q:
        like = f"%{q}%"
        query = query.filter(db.or_(
            Vehiculo.cc_chasis_marca.ilike(like),
            Vehiculo.cc_chasis_modelo.ilike(like),
            Vehiculo.div_marca.ilike(like),
            Vehiculo.div_modelo.ilike(like),
            Vehiculo.dorsal.ilike(like),
        ))

    vehiculos = query.order_by(Vehiculo.creado_en.desc()).paginate(page=page, per_page=25, error_out=False)
    return render_template("autoclub/vehiculos/lista.html",
                           vehiculos=vehiculos, q=q, tipo_vehiculo=tipo)


@autoclub_bp.route("/vehiculos/nuevo", methods=["GET", "POST"])
@login_required
def vehiculos_nuevo():
    pilotos = Piloto.query.filter_by(activo=True).order_by(Piloto.nombre.asc()).all()
    piloto_id_pre = request.args.get("piloto_id", type=int)
    if request.method == "POST":
        v = Vehiculo(
            tipo_vehiculo=request.form.get("tipo_vehiculo", "car_cross"),
            categoria=request.form.get("categoria", "").strip(),
            cc_chasis_marca=request.form.get("cc_chasis_marca", "").strip(),
            cc_chasis_modelo=request.form.get("cc_chasis_modelo", "").strip(),
            cc_num_chasis=request.form.get("cc_num_chasis", "").strip(),
            cc_pasaporte_rfeda=request.form.get("cc_pasaporte_rfeda", "").strip(),
            cc_pasaporte_autonomico=request.form.get("cc_pasaporte_autonomico", "").strip(),
            cc_num_homologacion=request.form.get("cc_num_homologacion", "").strip(),
            cc_motor_marca=request.form.get("cc_motor_marca", "").strip(),
            cc_motor_modelo=request.form.get("cc_motor_modelo", "").strip(),
            cc_motor_anio=request.form.get("cc_motor_anio", "").strip(),
            cc_cilindrada=_parse_int(request.form.get("cc_cilindrada")),
            cc_num_motor=request.form.get("cc_num_motor", "").strip(),
            div_marca=request.form.get("div_marca", "").strip(),
            div_modelo=request.form.get("div_modelo", "").strip(),
            div_num_bastidor=request.form.get("div_num_bastidor", "").strip(),
            div_pasaporte_rfeda=request.form.get("div_pasaporte_rfeda", "").strip(),
            div_pasaporte_autonomico=request.form.get("div_pasaporte_autonomico", "").strip(),
            div_homologacion=request.form.get("div_homologacion", "").strip(),
            div_motor_disposicion=request.form.get("div_motor_disposicion", "").strip(),
            div_motor_aspiracion=request.form.get("div_motor_aspiracion", "").strip(),
            div_cilindrada=_parse_int(request.form.get("div_cilindrada")),
            div_factor_correccion=_parse_float(request.form.get("div_factor_correccion")),
            div_traccion=request.form.get("div_traccion", "").strip(),
            arnes_marca=request.form.get("arnes_marca", "").strip(),
            arnes_homologacion=request.form.get("arnes_homologacion", "").strip(),
            arnes_caducidad=_parse_date(request.form.get("arnes_caducidad")),
            asiento_marca=request.form.get("asiento_marca", "").strip(),
            asiento_homologacion=request.form.get("asiento_homologacion", "").strip(),
            asiento_caducidad=_parse_date(request.form.get("asiento_caducidad")),
            extincion_marca=request.form.get("extincion_marca", "").strip(),
            extincion_tipo=request.form.get("extincion_tipo", "").strip(),
            extincion_revision=_parse_date(request.form.get("extincion_revision")),
            deposito_marca=request.form.get("deposito_marca", "").strip(),
            deposito_caducidad=_parse_date(request.form.get("deposito_caducidad")),
            red_estado=request.form.get("red_estado", "").strip(),
            red_homologacion=request.form.get("red_homologacion", "").strip(),
            dorsal=request.form.get("dorsal", "").strip(),
            num_transponder=request.form.get("num_transponder", "").strip(),
            piloto_id=_parse_int(request.form.get("piloto_id")),
        )
        db.session.add(v)
        db.session.commit()
        registrar_log("crear", "vehiculo", v.id, f"Vehículo creado: {v.nombre_display}")
        flash(f"Vehículo {v.nombre_display} creado correctamente.", "success")
        return redirect(url_for("autoclub.vehiculos_detalle", id=v.id))
    return render_template("autoclub/vehiculos/form.html", vehiculo=None,
                           pilotos=pilotos, piloto_id_pre=piloto_id_pre)


@autoclub_bp.route("/vehiculos/<int:id>")
@login_required
def vehiculos_detalle(id):
    vehiculo = Vehiculo.query.get_or_404(id)
    return render_template("autoclub/vehiculos/detalle.html", vehiculo=vehiculo)


@autoclub_bp.route("/vehiculos/<int:id>/editar", methods=["GET", "POST"])
@login_required
def vehiculos_editar(id):
    vehiculo = Vehiculo.query.get_or_404(id)
    pilotos = Piloto.query.filter_by(activo=True).order_by(Piloto.nombre.asc()).all()
    if request.method == "POST":
        vehiculo.tipo_vehiculo = request.form.get("tipo_vehiculo", "car_cross")
        vehiculo.categoria = request.form.get("categoria", "").strip()
        vehiculo.cc_chasis_marca = request.form.get("cc_chasis_marca", "").strip()
        vehiculo.cc_chasis_modelo = request.form.get("cc_chasis_modelo", "").strip()
        vehiculo.cc_num_chasis = request.form.get("cc_num_chasis", "").strip()
        vehiculo.cc_pasaporte_rfeda = request.form.get("cc_pasaporte_rfeda", "").strip()
        vehiculo.cc_pasaporte_autonomico = request.form.get("cc_pasaporte_autonomico", "").strip()
        vehiculo.cc_num_homologacion = request.form.get("cc_num_homologacion", "").strip()
        vehiculo.cc_motor_marca = request.form.get("cc_motor_marca", "").strip()
        vehiculo.cc_motor_modelo = request.form.get("cc_motor_modelo", "").strip()
        vehiculo.cc_motor_anio = request.form.get("cc_motor_anio", "").strip()
        vehiculo.cc_cilindrada = _parse_int(request.form.get("cc_cilindrada"))
        vehiculo.cc_num_motor = request.form.get("cc_num_motor", "").strip()
        vehiculo.div_marca = request.form.get("div_marca", "").strip()
        vehiculo.div_modelo = request.form.get("div_modelo", "").strip()
        vehiculo.div_num_bastidor = request.form.get("div_num_bastidor", "").strip()
        vehiculo.div_pasaporte_rfeda = request.form.get("div_pasaporte_rfeda", "").strip()
        vehiculo.div_pasaporte_autonomico = request.form.get("div_pasaporte_autonomico", "").strip()
        vehiculo.div_homologacion = request.form.get("div_homologacion", "").strip()
        vehiculo.div_motor_disposicion = request.form.get("div_motor_disposicion", "").strip()
        vehiculo.div_motor_aspiracion = request.form.get("div_motor_aspiracion", "").strip()
        vehiculo.div_cilindrada = _parse_int(request.form.get("div_cilindrada"))
        vehiculo.div_factor_correccion = _parse_float(request.form.get("div_factor_correccion"))
        vehiculo.div_traccion = request.form.get("div_traccion", "").strip()
        vehiculo.arnes_marca = request.form.get("arnes_marca", "").strip()
        vehiculo.arnes_homologacion = request.form.get("arnes_homologacion", "").strip()
        vehiculo.arnes_caducidad = _parse_date(request.form.get("arnes_caducidad"))
        vehiculo.asiento_marca = request.form.get("asiento_marca", "").strip()
        vehiculo.asiento_homologacion = request.form.get("asiento_homologacion", "").strip()
        vehiculo.asiento_caducidad = _parse_date(request.form.get("asiento_caducidad"))
        vehiculo.extincion_marca = request.form.get("extincion_marca", "").strip()
        vehiculo.extincion_tipo = request.form.get("extincion_tipo", "").strip()
        vehiculo.extincion_revision = _parse_date(request.form.get("extincion_revision"))
        vehiculo.deposito_marca = request.form.get("deposito_marca", "").strip()
        vehiculo.deposito_caducidad = _parse_date(request.form.get("deposito_caducidad"))
        vehiculo.red_estado = request.form.get("red_estado", "").strip()
        vehiculo.red_homologacion = request.form.get("red_homologacion", "").strip()
        vehiculo.dorsal = request.form.get("dorsal", "").strip()
        vehiculo.num_transponder = request.form.get("num_transponder", "").strip()
        vehiculo.piloto_id = _parse_int(request.form.get("piloto_id"))
        db.session.commit()
        registrar_log("editar", "vehiculo", vehiculo.id, f"Vehículo editado: {vehiculo.nombre_display}")
        flash("Vehículo actualizado correctamente.", "success")
        return redirect(url_for("autoclub.vehiculos_detalle", id=vehiculo.id))
    return render_template("autoclub/vehiculos/form.html", vehiculo=vehiculo, pilotos=pilotos, piloto_id_pre=None)


@autoclub_bp.route("/vehiculos/<int:id>/eliminar", methods=["POST"])
@login_required
def vehiculos_eliminar(id):
    vehiculo = Vehiculo.query.get_or_404(id)
    nombre = vehiculo.nombre_display
    db.session.delete(vehiculo)
    db.session.commit()
    registrar_log("eliminar", "vehiculo", id, f"Vehículo eliminado: {nombre}")
    flash(f"Vehículo {nombre} eliminado.", "info")
    return redirect(url_for("autoclub.vehiculos_lista"))


# ══════════════════════════════════════════════════════════════════════════════
#  INSCRIPCIONES
# ══════════════════════════════════════════════════════════════════════════════

@autoclub_bp.route("/inscripciones")
@login_required
def inscripciones_lista():
    q          = request.args.get("q", "").strip()
    estado     = request.args.get("estado", "")
    campeonato = request.args.get("campeonato", "")
    page       = request.args.get("page", 1, type=int)

    query = Inscripcion.query.join(Piloto)
    if estado:
        query = query.filter(Inscripcion.estado == estado)
    if campeonato:
        query = query.filter(Inscripcion.campeonato == campeonato)
    if q:
        like = f"%{q}%"
        query = query.filter(db.or_(
            Inscripcion.nombre_prueba.ilike(like),
            Piloto.nombre.ilike(like),
            Piloto.apellidos.ilike(like),
        ))

    inscripciones = query.order_by(Inscripcion.creado_en.desc()).paginate(page=page, per_page=25, error_out=False)
    campeonatos_disponibles = ["CEAX", "CERX", "FAG", "FADA", "FACM", "FACYL", "Autonómico", "Otro"]
    estados_disponibles = ["pendiente", "enviada", "aceptada", "rechazada"]
    return render_template("autoclub/inscripciones/lista.html",
                           inscripciones=inscripciones, q=q, estado=estado,
                           campeonato=campeonato,
                           campeonatos_disponibles=campeonatos_disponibles,
                           estados_disponibles=estados_disponibles)


@autoclub_bp.route("/inscripciones/nueva", methods=["GET", "POST"])
@login_required
def inscripciones_nueva():
    pilotos = Piloto.query.filter_by(activo=True).order_by(Piloto.nombre.asc()).all()
    vehiculos = Vehiculo.query.order_by(Vehiculo.creado_en.desc()).all()
    piloto_id_pre = request.args.get("piloto_id", type=int)
    fecha_pre = request.args.get("fecha", "")
    campeonatos = ["CEAX", "CERX", "FAG", "FADA", "FACM", "FACYL", "Autonómico", "Otro"]
    if request.method == "POST":
        insc = Inscripcion(
            nombre_prueba=request.form.get("nombre_prueba", "").strip(),
            fecha_prueba=_parse_date(request.form.get("fecha_prueba")),
            fecha_fin=_parse_date(request.form.get("fecha_fin")),
            fecha_plazo=_parse_date(request.form.get("fecha_plazo")),
            campeonato=request.form.get("campeonato", "").strip(),
            piloto_id=_parse_int(request.form.get("piloto_id")),
            vehiculo_id=_parse_int(request.form.get("vehiculo_id")),
            concursante=request.form.get("concursante", "").strip(),
            estado=request.form.get("estado", "pendiente"),
            notas=request.form.get("notas", "").strip(),
        )
        justificante = _save_doc(request.files.get("justificante"))
        if justificante:
            insc.justificante_filename = justificante
        db.session.add(insc)
        db.session.commit()
        registrar_log("crear", "inscripcion", insc.id, f"Inscripción creada: {insc.nombre_prueba}")
        flash(f"Inscripción para {insc.nombre_prueba} creada correctamente.", "success")
        return redirect(url_for("autoclub.inscripciones_detalle", id=insc.id))
    return render_template("autoclub/inscripciones/form.html", insc=None,
                           pilotos=pilotos, vehiculos=vehiculos,
                           piloto_id_pre=piloto_id_pre, fecha_pre=fecha_pre,
                           campeonatos=campeonatos)


@autoclub_bp.route("/inscripciones/<int:id>")
@login_required
def inscripciones_detalle(id):
    insc = Inscripcion.query.get_or_404(id)
    notas = AutoclubNota.query.filter_by(entidad="inscripcion", entidad_id=id).order_by(AutoclubNota.creado_en.desc()).all()
    gastos = GastoInscripcion.query.filter_by(inscripcion_id=id).order_by(GastoInscripcion.fecha.desc(), GastoInscripcion.creado_en.desc()).all()
    total_gastos = sum(g.importe for g in gastos)
    staff_asignado = insc.staff.all()
    staff_disponible = StaffMiembro.query.filter_by(activo=True).order_by(StaffMiembro.apellidos.asc()).all()
    staff_disponible = [s for s in staff_disponible if s not in staff_asignado]

    participantes = insc.participantes.order_by(InscripcionPiloto.id.asc()).all()
    ids_ya_inscritos = {insc.piloto_id} | {p.piloto_id for p in participantes}
    pilotos_disponibles = (
        Piloto.query.filter_by(activo=True)
        .filter(Piloto.id.notin_(ids_ya_inscritos))
        .order_by(Piloto.nombre.asc()).all()
    )
    vehiculos_todos = Vehiculo.query.order_by(Vehiculo.creado_en.desc()).all()

    return render_template(
        "autoclub/inscripciones/detalle.html",
        insc=insc,
        notas=notas,
        gastos=gastos,
        total_gastos=total_gastos,
        categorias_gasto=CATEGORIAS_GASTO,
        staff_asignado=staff_asignado,
        staff_disponible=staff_disponible,
        participantes=participantes,
        pilotos_disponibles=pilotos_disponibles,
        vehiculos_todos=vehiculos_todos,
        estados_inscripcion=ESTADOS_INSCRIPCION,
        today=date.today(),
    )


@autoclub_bp.route("/inscripciones/<int:id>/participantes/nuevo", methods=["POST"])
@login_required
def inscripcion_participante_nuevo(id):
    insc = Inscripcion.query.get_or_404(id)
    piloto_id = _parse_int(request.form.get("piloto_id"))
    if not piloto_id:
        flash("Selecciona un piloto para añadirlo a esta prueba.", "warning")
        return redirect(url_for("autoclub.inscripciones_detalle", id=id) + "#participantes")

    ya_inscrito = (piloto_id == insc.piloto_id) or insc.participantes.filter_by(piloto_id=piloto_id).first()
    if ya_inscrito:
        flash("Ese piloto ya está inscrito en esta prueba.", "warning")
        return redirect(url_for("autoclub.inscripciones_detalle", id=id) + "#participantes")

    participante = InscripcionPiloto(
        inscripcion_id=insc.id,
        piloto_id=piloto_id,
        vehiculo_id=_parse_int(request.form.get("vehiculo_id")),
        concursante=request.form.get("concursante", "").strip(),
    )
    db.session.add(participante)
    db.session.commit()
    registrar_log("editar", "inscripcion", insc.id,
                  f"Piloto añadido a la inscripción: {insc.nombre_prueba} — {participante.piloto.nombre_completo}")
    flash("Piloto añadido a la prueba.", "success")
    return redirect(url_for("autoclub.inscripciones_detalle", id=id) + "#participantes")


@autoclub_bp.route("/inscripciones/<int:id>/participantes/<int:pid>/eliminar", methods=["POST"])
@login_required
def inscripcion_participante_eliminar(id, pid):
    participante = InscripcionPiloto.query.filter_by(id=pid, inscripcion_id=id).first_or_404()
    nombre = participante.piloto.nombre_completo if participante.piloto else "piloto"
    db.session.delete(participante)
    db.session.commit()
    registrar_log("editar", "inscripcion", id, f"Piloto quitado de la inscripción: {nombre}")
    flash("Piloto quitado de la prueba.", "info")
    return redirect(url_for("autoclub.inscripciones_detalle", id=id) + "#participantes")


@autoclub_bp.route("/inscripciones/<int:id>/estado-principal", methods=["POST"])
@login_required
def inscripcion_estado_principal(id):
    """Cambia el estado individual del piloto principal (Inscripcion.estado)
    sin pasar por el formulario completo de edición."""
    insc = Inscripcion.query.get_or_404(id)
    nuevo_estado = request.form.get("estado", "")
    if nuevo_estado not in ESTADOS_INSCRIPCION:
        flash(f"Estado '{nuevo_estado}' no válido.", "warning")
        return redirect(url_for("autoclub.inscripciones_detalle", id=id) + "#participantes")
    insc.estado = nuevo_estado
    db.session.commit()
    registrar_log("editar", "inscripcion", id,
                  f"Estado del piloto principal cambiado a {nuevo_estado}: {insc.nombre_prueba}")
    flash(f"Estado de {insc.piloto.nombre_completo if insc.piloto else 'piloto principal'} actualizado a {nuevo_estado}.", "success")
    return redirect(url_for("autoclub.inscripciones_detalle", id=id) + "#participantes")


@autoclub_bp.route("/inscripciones/<int:id>/participantes/<int:pid>/estado", methods=["POST"])
@login_required
def inscripcion_participante_estado(id, pid):
    """Cambia el estado individual de un piloto adicional."""
    participante = InscripcionPiloto.query.filter_by(id=pid, inscripcion_id=id).first_or_404()
    nuevo_estado = request.form.get("estado", "")
    if nuevo_estado not in ESTADOS_INSCRIPCION:
        flash(f"Estado '{nuevo_estado}' no válido.", "warning")
        return redirect(url_for("autoclub.inscripciones_detalle", id=id) + "#participantes")
    participante.estado = nuevo_estado
    db.session.commit()
    nombre = participante.piloto.nombre_completo if participante.piloto else "piloto"
    registrar_log("editar", "inscripcion", id, f"Estado de {nombre} cambiado a {nuevo_estado}")
    flash(f"Estado de {nombre} actualizado a {nuevo_estado}.", "success")
    return redirect(url_for("autoclub.inscripciones_detalle", id=id) + "#participantes")


@autoclub_bp.route("/inscripciones/<int:id>/editar", methods=["GET", "POST"])
@login_required
def inscripciones_editar(id):
    insc = Inscripcion.query.get_or_404(id)
    pilotos = Piloto.query.filter_by(activo=True).order_by(Piloto.nombre.asc()).all()
    vehiculos = Vehiculo.query.order_by(Vehiculo.creado_en.desc()).all()
    campeonatos = ["CEAX", "CERX", "FAG", "FADA", "FACM", "FACYL", "Autonómico", "Otro"]
    if request.method == "POST":
        insc.nombre_prueba = request.form.get("nombre_prueba", "").strip()
        insc.fecha_prueba = _parse_date(request.form.get("fecha_prueba"))
        insc.fecha_fin = _parse_date(request.form.get("fecha_fin"))
        insc.fecha_plazo = _parse_date(request.form.get("fecha_plazo"))
        insc.campeonato = request.form.get("campeonato", "").strip()
        insc.piloto_id = _parse_int(request.form.get("piloto_id"))
        insc.vehiculo_id = _parse_int(request.form.get("vehiculo_id"))
        insc.concursante = request.form.get("concursante", "").strip()
        insc.estado = request.form.get("estado", "pendiente")
        insc.notas = request.form.get("notas", "").strip()
        justificante = _save_doc(request.files.get("justificante"))
        if justificante:
            _delete_doc(insc.justificante_filename)
            insc.justificante_filename = justificante
        db.session.commit()
        registrar_log("editar", "inscripcion", insc.id, f"Inscripción editada: {insc.nombre_prueba}")
        flash("Inscripción actualizada correctamente.", "success")
        return redirect(url_for("autoclub.inscripciones_detalle", id=insc.id))
    return render_template("autoclub/inscripciones/form.html", insc=insc,
                           pilotos=pilotos, vehiculos=vehiculos,
                           piloto_id_pre=None, campeonatos=campeonatos)


@autoclub_bp.route("/inscripciones/<int:id>/gastos/exportar")
@login_required
def inscripciones_gastos_exportar(id):
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    insc   = Inscripcion.query.get_or_404(id)
    gastos = GastoInscripcion.query.filter_by(inscripcion_id=id).order_by(
        GastoInscripcion.fecha.asc(), GastoInscripcion.creado_en.asc()
    ).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Gastos"

    # ── Paleta ───────────────────────────────────────────────
    NEGRO         = "111111"
    VERDE         = "16A34A"
    GRIS_CABECERA = "F3F4F6"
    GRIS_BORDE    = "D1D5DB"

    def fill(hex_color):
        return PatternFill("solid", fgColor=hex_color)

    def borde_fino():
        lado = Side(style="thin", color=GRIS_BORDE)
        return Border(left=lado, right=lado, top=lado, bottom=lado)

    # ── Bloque de cabecera (filas 1-5) ───────────────────────
    info = [
        ("Prueba",       insc.nombre_prueba),
        ("Piloto",       insc.piloto.nombre_completo if insc.piloto else ""),
        ("Fecha prueba", insc.rango_fechas),
        ("Estado",       insc.estado_grupo.capitalize()),
        ("Exportado",    __import__("datetime").date.today().strftime("%d/%m/%Y")),
    ]
    for r, (etiqueta, valor) in enumerate(info, start=1):
        ws.cell(r, 1, etiqueta).font  = Font(bold=True, color=NEGRO, size=9)
        ws.cell(r, 1).fill            = fill(GRIS_CABECERA)
        ws.cell(r, 2, valor).font     = Font(color=NEGRO, size=9)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)

    # Fila separadora vacía
    ws.append([])

    # ── Cabecera de tabla (fila 7) ───────────────────────────
    HDR_ROW = 7
    headers = ["Concepto", "Categoría", "Fecha", "Importe (€)"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(HDR_ROW, c, h)
        cell.font      = Font(bold=True, color="FFFFFF", size=10)
        cell.fill      = fill(NEGRO)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = borde_fino()
    ws.row_dimensions[HDR_ROW].height = 22

    # ── Filas de datos ───────────────────────────────────────
    cat_colors = {
        "Inscripción": "DBEAFE", "Combustible": "FEF3C7",
        "Alojamiento": "F3E8FF", "Dietas":      "FCE7F3",
        "Transporte":  "E0F2FE", "Repuestos":   "FEE2E2",
        "Neumáticos":  "FFEDD5", "Equipación":  "ECFDF5",
        "Otros":       "F9FAFB",
    }
    total = 0.0
    for i, g in enumerate(gastos):
        row = HDR_ROW + 1 + i
        bg  = cat_colors.get(g.categoria or "Otros", "F9FAFB")
        zebra = "FFFFFF" if i % 2 == 0 else "F9FAFB"

        concepto_cell = ws.cell(row, 1, g.concepto)
        concepto_cell.font      = Font(size=9)
        concepto_cell.fill      = fill(zebra)
        concepto_cell.border    = borde_fino()

        cat_cell = ws.cell(row, 2, g.categoria or "")
        cat_cell.font      = Font(size=9, color=NEGRO)
        cat_cell.fill      = fill(bg)
        cat_cell.alignment = Alignment(horizontal="center")
        cat_cell.border    = borde_fino()

        fecha_cell = ws.cell(row, 3, g.fecha.strftime("%d/%m/%Y") if g.fecha else "")
        fecha_cell.font      = Font(size=9)
        fecha_cell.fill      = fill(zebra)
        fecha_cell.alignment = Alignment(horizontal="center")
        fecha_cell.border    = borde_fino()

        imp_cell = ws.cell(row, 4, g.importe)
        imp_cell.font           = Font(size=9)
        imp_cell.fill           = fill(zebra)
        imp_cell.number_format  = '#,##0.00 "€"'
        imp_cell.alignment      = Alignment(horizontal="right")
        imp_cell.border         = borde_fino()
        total += g.importe

    # ── Fila de total ────────────────────────────────────────
    total_row = HDR_ROW + 1 + len(gastos) + 1
    ws.cell(total_row, 1, "TOTAL").font      = Font(bold=True, color="FFFFFF", size=10)
    ws.cell(total_row, 1).fill               = fill(VERDE)
    ws.cell(total_row, 1).alignment          = Alignment(horizontal="right")
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=3)
    imp_total = ws.cell(total_row, 4, total)
    imp_total.font          = Font(bold=True, color="FFFFFF", size=10)
    imp_total.fill          = fill(VERDE)
    imp_total.number_format = '#,##0.00 "€"'
    imp_total.alignment     = Alignment(horizontal="right")

    # ── Anchos de columna ────────────────────────────────────
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 16

    # ── Inmovilizar primera fila de tabla ────────────────────
    ws.freeze_panes = ws.cell(HDR_ROW + 1, 1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    nombre = f"gastos_{insc.nombre_prueba[:40].replace(' ', '_')}.xlsx"
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
    )


@autoclub_bp.route("/inscripciones/<int:id>/eliminar", methods=["POST"])
@login_required
def inscripciones_eliminar(id):
    insc = Inscripcion.query.get_or_404(id)
    nombre = insc.nombre_prueba
    _delete_doc(insc.justificante_filename)
    db.session.delete(insc)
    db.session.commit()
    registrar_log("eliminar", "inscripcion", id, f"Inscripción eliminada: {nombre}")
    flash(f"Inscripción {nombre} eliminada.", "info")
    return redirect(url_for("autoclub.inscripciones_lista"))


@autoclub_bp.route("/inscripciones/<int:id>/estado", methods=["POST"])
@login_required
def inscripciones_estado(id):
    insc = Inscripcion.query.get_or_404(id)
    nuevo_estado = request.form.get("estado", "pendiente")
    insc.estado = nuevo_estado
    db.session.commit()
    registrar_log("cambiar_estado", "inscripcion", insc.id,
                  f"Estado inscripción cambiado a {nuevo_estado}: {insc.nombre_prueba}")
    flash(f"Estado actualizado a {nuevo_estado}.", "success")
    return redirect(url_for("autoclub.inscripciones_detalle", id=insc.id))


@autoclub_bp.route("/inscripciones/<int:id>/notas/nueva", methods=["POST"])
@login_required
def inscripcion_nota_nueva(id):
    insc = Inscripcion.query.get_or_404(id)
    contenido = request.form.get("contenido", "").strip()
    if contenido:
        nota = AutoclubNota(
            entidad="inscripcion",
            entidad_id=insc.id,
            entidad_nombre=insc.nombre_prueba,
            contenido=contenido,
            usuario_nombre=current_user.username if hasattr(current_user, "username") else "",
        )
        db.session.add(nota)
        db.session.commit()
        flash("Nota añadida.", "success")
    return redirect(url_for("autoclub.inscripciones_detalle", id=id) + "#notas")


@autoclub_bp.route("/inscripciones/<int:id>/notas/<int:nid>/eliminar", methods=["POST"])
@login_required
def inscripcion_nota_eliminar(id, nid):
    nota = AutoclubNota.query.get_or_404(nid)
    db.session.delete(nota)
    db.session.commit()
    flash("Nota eliminada.", "info")
    return redirect(url_for("autoclub.inscripciones_detalle", id=id) + "#notas")


@autoclub_bp.route("/inscripciones/<int:id>/gastos/nuevo", methods=["POST"])
@login_required
def inscripcion_gasto_nuevo(id):
    insc = Inscripcion.query.get_or_404(id)
    concepto = request.form.get("concepto", "").strip()
    importe_raw = request.form.get("importe", "0").replace(",", ".").strip()
    try:
        importe = float(importe_raw)
    except ValueError:
        importe = 0.0
    categoria = request.form.get("categoria", "Otros")
    fecha_raw = request.form.get("fecha", "").strip()
    fecha = None
    if fecha_raw:
        try:
            fecha = date.fromisoformat(fecha_raw)
        except ValueError:
            pass
    if concepto:
        comprobante_fn = _leer_doc_gasto("comprobante", "comprobante")
        factura_fn     = _leer_doc_gasto("factura", "factura")

        gasto = GastoInscripcion(
            inscripcion_id=insc.id,
            concepto=concepto,
            importe=importe,
            categoria=categoria,
            fecha=fecha,
            documento_filename=comprobante_fn,
            factura_filename=factura_fn,
        )
        db.session.add(gasto)
        db.session.commit()
        flash("Gasto añadido.", "success")
    return redirect(url_for("autoclub.inscripciones_detalle", id=id) + "#gastos")


@autoclub_bp.route("/inscripciones/<int:id>/gastos/<int:gid>/documento/<tipo>", methods=["POST"])
@login_required
def inscripcion_gasto_documento(id, gid, tipo):
    if tipo not in ("comprobante", "factura"):
        abort(404)
    gasto = GastoInscripcion.query.get_or_404(gid)
    campo = "documento_filename" if tipo == "comprobante" else "factura_filename"
    doc_file = request.files.get(tipo)
    if doc_file and doc_file.filename:
        doc_fn = _save_doc(doc_file)
        if doc_fn:
            anterior = getattr(gasto, campo)
            if anterior:
                _delete_doc(anterior)
            setattr(gasto, campo, doc_fn)
            db.session.commit()
            flash(f"{tipo.capitalize()} adjuntado.", "success")
        else:
            flash(f"El/la {tipo} «{doc_file.filename}» no se pudo adjuntar: formato no admitido "
                  f"(usa PDF, JPG, PNG, WEBP, DOC o DOCX).", "warning")
    return redirect(url_for("autoclub.inscripciones_detalle", id=id) + "#gastos")


@autoclub_bp.route("/inscripciones/<int:id>/gastos/<int:gid>/eliminar", methods=["POST"])
@login_required
def inscripcion_gasto_eliminar(id, gid):
    gasto = GastoInscripcion.query.get_or_404(gid)
    if gasto.documento_filename:
        _delete_doc(gasto.documento_filename)
    if gasto.factura_filename:
        _delete_doc(gasto.factura_filename)
    db.session.delete(gasto)
    db.session.commit()
    flash("Gasto eliminado.", "info")
    return redirect(url_for("autoclub.inscripciones_detalle", id=id) + "#gastos")


# ══════════════════════════════════════════════════════════════════════════════
#  GASTOS (sección global — CRUD, vinculables a una inscripción o a un evento)
# ══════════════════════════════════════════════════════════════════════════════

def _gasto_origen_url(gasto):
    """Endpoint al que volver tras crear/editar/eliminar un gasto: la ficha
    de su inscripción o evento de origen."""
    if gasto.inscripcion_id:
        return url_for("autoclub.gastos_origen", tipo="inscripcion", id=gasto.inscripcion_id)
    if gasto.evento_id:
        return url_for("autoclub.gastos_origen", tipo="evento", id=gasto.evento_id)
    return url_for("autoclub.gastos_lista")


@autoclub_bp.route("/gastos")
@login_required
def gastos_lista():
    """Página principal de Gastos: inscripciones y eventos con su nº de
    gastos y total acumulado, para entrar y ver/añadir/editar/borrar los
    gastos de cada uno. Incluye también las que todavía no tienen gastos."""
    from sqlalchemy import func
    from sqlalchemy.orm import joinedload
    from app.models.competicion_evento import CompeticionEvento

    conteo_insc = dict(
        db.session.query(GastoInscripcion.inscripcion_id, func.count(GastoInscripcion.id))
        .filter(GastoInscripcion.inscripcion_id.isnot(None))
        .group_by(GastoInscripcion.inscripcion_id).all()
    )
    total_insc = dict(
        db.session.query(GastoInscripcion.inscripcion_id, func.coalesce(func.sum(GastoInscripcion.importe), 0.0))
        .filter(GastoInscripcion.inscripcion_id.isnot(None))
        .group_by(GastoInscripcion.inscripcion_id).all()
    )
    conteo_ev = dict(
        db.session.query(GastoInscripcion.evento_id, func.count(GastoInscripcion.id))
        .filter(GastoInscripcion.evento_id.isnot(None))
        .group_by(GastoInscripcion.evento_id).all()
    )
    total_ev = dict(
        db.session.query(GastoInscripcion.evento_id, func.coalesce(func.sum(GastoInscripcion.importe), 0.0))
        .filter(GastoInscripcion.evento_id.isnot(None))
        .group_by(GastoInscripcion.evento_id).all()
    )

    inscripciones = (
        Inscripcion.query.options(joinedload(Inscripcion.piloto))
        .order_by(Inscripcion.fecha_prueba.desc().nullslast()).all()
    )
    eventos = CompeticionEvento.query.order_by(CompeticionEvento.fecha.desc()).all()

    filas_insc = [
        {"obj": i, "n": conteo_insc.get(i.id, 0), "total": total_insc.get(i.id, 0.0)}
        for i in inscripciones
    ]
    filas_ev = [
        {"obj": e, "n": conteo_ev.get(e.id, 0), "total": total_ev.get(e.id, 0.0)}
        for e in eventos
    ]

    total_general = sum(total_insc.values()) + sum(total_ev.values())

    return render_template(
        "autoclub/gastos/lista.html",
        filas_insc=filas_insc, filas_ev=filas_ev, total_general=total_general,
    )


@autoclub_bp.route("/gastos/todos")
@login_required
def gastos_todos():
    """Vista plana con todos los gastos juntos, con filtros — la opción
    "ver todo" para quien prefiera una tabla en vez de navegar por origen."""
    from sqlalchemy.orm import joinedload
    from app.models.competicion_evento import CompeticionEvento

    origen    = request.args.get("origen", "")       # "" | "inscripcion" | "evento"
    categoria = request.args.get("categoria", "")
    q         = request.args.get("q", "").strip()

    query = GastoInscripcion.query.options(
        joinedload(GastoInscripcion.inscripcion).joinedload(Inscripcion.piloto),
        joinedload(GastoInscripcion.evento),
    )
    if origen == "inscripcion":
        query = query.filter(GastoInscripcion.inscripcion_id.isnot(None))
    elif origen == "evento":
        query = query.filter(GastoInscripcion.evento_id.isnot(None))
    if categoria:
        query = query.filter(GastoInscripcion.categoria == categoria)
    if q:
        query = query.filter(GastoInscripcion.concepto.ilike(f"%{q}%"))

    gastos = query.order_by(
        GastoInscripcion.fecha.desc().nullslast(), GastoInscripcion.creado_en.desc()
    ).all()
    total = sum(g.importe for g in gastos)

    return render_template(
        "autoclub/gastos/todos.html",
        gastos=gastos, total=total, categorias_gasto=CATEGORIAS_GASTO,
        origen=origen, categoria=categoria, q=q,
    )


@autoclub_bp.route("/gastos/<tipo>/<int:id>")
@login_required
def gastos_origen(tipo, id):
    """Todos los gastos de UNA inscripción o UN evento en concreto, con
    alta/edición/borrado propios (sin tener que volver a elegir el origen)."""
    from sqlalchemy.orm import joinedload
    from app.models.competicion_evento import CompeticionEvento
    if tipo not in ("inscripcion", "evento"):
        abort(404)

    if tipo == "inscripcion":
        origen_obj = Inscripcion.query.options(joinedload(Inscripcion.piloto)).get_or_404(id)
        gastos = (
            GastoInscripcion.query.filter_by(inscripcion_id=id)
            .order_by(GastoInscripcion.fecha.desc().nullslast(), GastoInscripcion.creado_en.desc())
            .all()
        )
    else:
        origen_obj = CompeticionEvento.query.get_or_404(id)
        gastos = (
            GastoInscripcion.query.filter_by(evento_id=id)
            .order_by(GastoInscripcion.fecha.desc().nullslast(), GastoInscripcion.creado_en.desc())
            .all()
        )

    total = sum(g.importe for g in gastos)
    return render_template(
        "autoclub/gastos/origen.html",
        tipo=tipo, origen_obj=origen_obj, gastos=gastos, total=total,
        categorias_gasto=CATEGORIAS_GASTO,
    )


@autoclub_bp.route("/gastos/nuevo", methods=["GET", "POST"])
@login_required
def gastos_nuevo():
    from app.models.competicion_evento import CompeticionEvento
    inscripciones = Inscripcion.query.order_by(Inscripcion.fecha_prueba.desc()).all()
    eventos       = CompeticionEvento.query.order_by(CompeticionEvento.fecha.desc()).all()

    # Permite pre-seleccionar el origen al llegar desde la ficha de una
    # inscripción/evento concretos (?tipo=evento&id=3).
    tipo_pre = request.args.get("tipo", "")
    id_pre   = request.args.get("id", type=int)

    if request.method == "POST":
        tipo_origen    = request.form.get("tipo_origen", "inscripcion")
        inscripcion_id = request.form.get("inscripcion_id", type=int) if tipo_origen == "inscripcion" else None
        evento_id      = request.form.get("evento_id", type=int) if tipo_origen == "evento" else None

        concepto = request.form.get("concepto", "").strip()
        importe_raw = request.form.get("importe", "0").replace(",", ".").strip()
        try:
            importe = float(importe_raw)
        except ValueError:
            importe = 0.0
        categoria = request.form.get("categoria", "Otros")
        fecha_raw = request.form.get("fecha", "").strip()
        fecha = None
        if fecha_raw:
            try:
                fecha = date.fromisoformat(fecha_raw)
            except ValueError:
                pass

        if not concepto:
            flash("El concepto es obligatorio.", "warning")
        elif not inscripcion_id and not evento_id:
            flash("Selecciona una inscripción o un evento al que vincular el gasto.", "warning")
        else:
            comprobante_fn = _leer_doc_gasto("comprobante", "comprobante")
            factura_fn     = _leer_doc_gasto("factura", "factura")
            gasto = GastoInscripcion(
                inscripcion_id=inscripcion_id,
                evento_id=evento_id,
                concepto=concepto,
                importe=importe,
                categoria=categoria,
                fecha=fecha,
                documento_filename=comprobante_fn,
                factura_filename=factura_fn,
            )
            db.session.add(gasto)
            db.session.commit()
            flash("Gasto añadido.", "success")
            return redirect(_gasto_origen_url(gasto))

    return render_template(
        "autoclub/gastos/form.html", gasto=None,
        inscripciones=inscripciones, eventos=eventos, categorias_gasto=CATEGORIAS_GASTO,
        tipo_pre=tipo_pre, id_pre=id_pre, today=date.today(),
    )


@autoclub_bp.route("/gastos/<int:id>/editar", methods=["GET", "POST"])
@login_required
def gastos_editar(id):
    from app.models.competicion_evento import CompeticionEvento
    gasto = GastoInscripcion.query.get_or_404(id)
    inscripciones = Inscripcion.query.order_by(Inscripcion.fecha_prueba.desc()).all()
    eventos       = CompeticionEvento.query.order_by(CompeticionEvento.fecha.desc()).all()

    if request.method == "POST":
        concepto = request.form.get("concepto", "").strip()
        if not concepto:
            flash("El concepto es obligatorio.", "warning")
            return render_template(
                "autoclub/gastos/form.html", gasto=gasto,
                inscripciones=inscripciones, eventos=eventos, categorias_gasto=CATEGORIAS_GASTO,
                tipo_pre="", id_pre=None, today=date.today(),
            )

        tipo_origen = request.form.get("tipo_origen", "inscripcion")
        if tipo_origen == "inscripcion":
            gasto.inscripcion_id = request.form.get("inscripcion_id", type=int)
            gasto.evento_id = None
        else:
            gasto.evento_id = request.form.get("evento_id", type=int)
            gasto.inscripcion_id = None

        gasto.concepto = concepto
        importe_raw = request.form.get("importe", "0").replace(",", ".").strip()
        try:
            gasto.importe = float(importe_raw)
        except ValueError:
            gasto.importe = 0.0
        gasto.categoria = request.form.get("categoria", "Otros")
        fecha_raw = request.form.get("fecha", "").strip()
        gasto.fecha = None
        if fecha_raw:
            try:
                gasto.fecha = date.fromisoformat(fecha_raw)
            except ValueError:
                pass

        comprobante_fn = _leer_doc_gasto("comprobante", "comprobante")
        if comprobante_fn:
            if gasto.documento_filename:
                _delete_doc(gasto.documento_filename)
            gasto.documento_filename = comprobante_fn
        factura_fn = _leer_doc_gasto("factura", "factura")
        if factura_fn:
            if gasto.factura_filename:
                _delete_doc(gasto.factura_filename)
            gasto.factura_filename = factura_fn

        db.session.commit()
        flash("Gasto actualizado.", "success")
        return redirect(_gasto_origen_url(gasto))

    return render_template(
        "autoclub/gastos/form.html", gasto=gasto,
        inscripciones=inscripciones, eventos=eventos, categorias_gasto=CATEGORIAS_GASTO,
        tipo_pre="", id_pre=None, today=date.today(),
    )


@autoclub_bp.route("/gastos/<int:id>/eliminar", methods=["POST"])
@login_required
def gastos_eliminar(id):
    gasto = GastoInscripcion.query.get_or_404(id)
    volver_a = _gasto_origen_url(gasto)
    if gasto.documento_filename:
        _delete_doc(gasto.documento_filename)
    if gasto.factura_filename:
        _delete_doc(gasto.factura_filename)
    db.session.delete(gasto)
    db.session.commit()
    flash("Gasto eliminado.", "info")
    return redirect(volver_a)


@autoclub_bp.route("/gastos/<int:id>/documento/<tipo>", methods=["POST"])
@login_required
def gastos_documento(id, tipo):
    if tipo not in ("comprobante", "factura"):
        abort(404)
    gasto = GastoInscripcion.query.get_or_404(id)
    campo = "documento_filename" if tipo == "comprobante" else "factura_filename"
    doc_file = request.files.get(tipo)
    if doc_file and doc_file.filename:
        doc_fn = _save_doc(doc_file)
        if doc_fn:
            anterior = getattr(gasto, campo)
            if anterior:
                _delete_doc(anterior)
            setattr(gasto, campo, doc_fn)
            db.session.commit()
            flash(f"{tipo.capitalize()} adjuntado.", "success")
        else:
            flash(f"El/la {tipo} «{doc_file.filename}» no se pudo adjuntar: formato no admitido "
                  f"(usa PDF, JPG, PNG, WEBP, DOC o DOCX).", "warning")
    return redirect(_gasto_origen_url(gasto))


@autoclub_bp.route("/inscripciones/<int:id>/resultado/guardar", methods=["POST"])
@login_required
def inscripcion_resultado_guardar(id):
    insc = Inscripcion.query.get_or_404(id)
    r = insc.resultado or ResultadoInscripcion(inscripcion_id=insc.id)

    def _int(k):
        v = request.form.get(k, "").strip()
        try: return int(v) if v else None
        except ValueError: return None

    def _float(k):
        v = request.form.get(k, "").replace(",", ".").strip()
        try: return float(v) if v else None
        except ValueError: return None

    r.posicion_salida     = _int("posicion_salida")
    r.posicion_final      = _int("posicion_final")
    r.mejor_vuelta        = request.form.get("mejor_vuelta", "").strip() or None
    r.tiempo_total        = request.form.get("tiempo_total", "").strip() or None
    r.vueltas_completadas = _int("vueltas_completadas")
    r.puntos              = _float("puntos")
    r.abandono            = bool(request.form.get("abandono"))
    r.motivo_abandono     = request.form.get("motivo_abandono", "").strip() or None
    r.observaciones       = request.form.get("observaciones", "").strip() or None

    if not r.id:
        db.session.add(r)
    db.session.commit()
    flash("Resultado guardado.", "success")
    return redirect(url_for("autoclub.inscripciones_detalle", id=id) + "#resultado")


@autoclub_bp.route("/inscripciones/<int:id>/resultado/eliminar", methods=["POST"])
@login_required
def inscripcion_resultado_eliminar(id):
    insc = Inscripcion.query.get_or_404(id)
    if insc.resultado:
        db.session.delete(insc.resultado)
        db.session.commit()
        flash("Resultado eliminado.", "info")
    return redirect(url_for("autoclub.inscripciones_detalle", id=id) + "#resultado")


# ══════════════════════════════════════════════════════════════════════════════
#  STAFF
# ══════════════════════════════════════════════════════════════════════════════

@autoclub_bp.route("/staff")
@login_required
def staff_lista():
    q = request.args.get("q", "").strip()
    tipo = request.args.get("tipo", "").strip()
    page = request.args.get("page", 1, type=int)
    query = StaffMiembro.query
    if q:
        query = query.filter(
            db.or_(
                StaffMiembro.nombre.ilike(f"%{q}%"),
                StaffMiembro.apellidos.ilike(f"%{q}%"),
                StaffMiembro.email.ilike(f"%{q}%"),
            )
        )
    if tipo:
        query = query.filter(StaffMiembro.tipo == tipo)
    miembros = query.order_by(StaffMiembro.apellidos.asc(), StaffMiembro.nombre.asc()).paginate(page=page, per_page=30, error_out=False)
    return render_template("autoclub/staff/lista.html", miembros=miembros, q=q, tipo=tipo, tipos=TIPOS_STAFF)


@autoclub_bp.route("/staff/nuevo", methods=["GET", "POST"])
@login_required
def staff_nuevo():
    if request.method == "POST":
        m = StaffMiembro(
            nombre=request.form.get("nombre", "").strip(),
            apellidos=request.form.get("apellidos", "").strip(),
            tipo=request.form.get("tipo", "Mecánico"),
            telefono=request.form.get("telefono", "").strip(),
            email=request.form.get("email", "").strip(),
            dni=request.form.get("dni", "").strip(),
            notas=request.form.get("notas", "").strip(),
        )
        db.session.add(m)
        db.session.commit()
        registrar_log("crear", "staff", m.id, f"Staff creado: {m.nombre_completo}")
        flash(f"{m.nombre_completo} añadido al staff.", "success")
        return redirect(url_for("autoclub.staff_detalle", id=m.id))
    return render_template("autoclub/staff/form.html", miembro=None, tipos=TIPOS_STAFF)


@autoclub_bp.route("/staff/<int:id>")
@login_required
def staff_detalle(id):
    m = StaffMiembro.query.get_or_404(id)
    inscripciones = m.inscripciones
    return render_template("autoclub/staff/detalle.html", m=m, inscripciones=inscripciones)


@autoclub_bp.route("/staff/<int:id>/editar", methods=["GET", "POST"])
@login_required
def staff_editar(id):
    m = StaffMiembro.query.get_or_404(id)
    if request.method == "POST":
        m.nombre   = request.form.get("nombre", "").strip()
        m.apellidos = request.form.get("apellidos", "").strip()
        m.tipo     = request.form.get("tipo", m.tipo)
        m.telefono = request.form.get("telefono", "").strip()
        m.email    = request.form.get("email", "").strip()
        m.dni      = request.form.get("dni", "").strip()
        m.notas    = request.form.get("notas", "").strip()
        db.session.commit()
        registrar_log("editar", "staff", m.id, f"Staff editado: {m.nombre_completo}")
        flash("Datos actualizados.", "success")
        return redirect(url_for("autoclub.staff_detalle", id=m.id))
    return render_template("autoclub/staff/form.html", miembro=m, tipos=TIPOS_STAFF)


@autoclub_bp.route("/staff/<int:id>/toggle", methods=["POST"])
@login_required
def staff_toggle(id):
    m = StaffMiembro.query.get_or_404(id)
    m.activo = not m.activo
    db.session.commit()
    flash(f"{m.nombre_completo} {'activado' if m.activo else 'desactivado'}.", "info")
    return redirect(url_for("autoclub.staff_detalle", id=m.id))


@autoclub_bp.route("/staff/<int:id>/eliminar", methods=["POST"])
@login_required
def staff_eliminar(id):
    m = StaffMiembro.query.get_or_404(id)
    nombre = m.nombre_completo
    db.session.delete(m)
    db.session.commit()
    registrar_log("eliminar", "staff", id, f"Staff eliminado: {nombre}")
    flash(f"{nombre} eliminado.", "info")
    return redirect(url_for("autoclub.staff_lista"))


@autoclub_bp.route("/inscripciones/<int:id>/staff/añadir", methods=["POST"])
@login_required
def inscripcion_staff_añadir(id):
    insc = Inscripcion.query.get_or_404(id)
    staff_id = request.form.get("staff_id", type=int)
    if staff_id:
        miembro = StaffMiembro.query.get(staff_id)
        if miembro and miembro not in insc.staff.all():
            insc.staff.append(miembro)
            db.session.commit()
            flash(f"{miembro.nombre_completo} añadido a la inscripción.", "success")
    return redirect(url_for("autoclub.inscripciones_detalle", id=id) + "#staff")


@autoclub_bp.route("/inscripciones/<int:id>/staff/<int:sid>/quitar", methods=["POST"])
@login_required
def inscripcion_staff_quitar(id, sid):
    insc = Inscripcion.query.get_or_404(id)
    miembro = StaffMiembro.query.get_or_404(sid)
    if miembro in insc.staff.all():
        insc.staff.remove(miembro)
        db.session.commit()
        flash(f"{miembro.nombre_completo} quitado de la inscripción.", "info")
    return redirect(url_for("autoclub.inscripciones_detalle", id=id) + "#staff")


# ══════════════════════════════════════════════════════════════════════════════
#  NOTAS
# ══════════════════════════════════════════════════════════════════════════════

@autoclub_bp.route("/notas")
@login_required
def notas_lista():
    q = request.args.get("q", "").strip()
    page = request.args.get("page", 1, type=int)
    query = AutoclubNota.query
    if q:
        query = query.filter(
            db.or_(
                AutoclubNota.entidad_nombre.ilike(f"%{q}%"),
                AutoclubNota.contenido.ilike(f"%{q}%"),
            )
        )
    notas = query.order_by(AutoclubNota.creado_en.desc()).paginate(page=page, per_page=30, error_out=False)
    return render_template("autoclub/notas/lista.html", notas=notas, q=q)


@autoclub_bp.route("/socios/<int:id>/notas/nueva", methods=["POST"])
@login_required
def socio_nota_nueva(id):
    s = Socio.query.get_or_404(id)
    contenido = request.form.get("contenido", "").strip()
    if contenido:
        nota = AutoclubNota(
            entidad="socio",
            entidad_id=s.id,
            entidad_nombre=f"{s.nombre} {s.apellidos or ''}".strip(),
            contenido=contenido,
            usuario_nombre=current_user.username if hasattr(current_user, "username") else "",
        )
        db.session.add(nota)
        db.session.commit()
        flash("Nota añadida.", "success")
    return_to = request.form.get("return_to", "detalle")
    if return_to == "lista":
        return redirect(url_for("autoclub.lista"))
    return redirect(url_for("autoclub.detalle", id=id) + "#notas")


@autoclub_bp.route("/socios/<int:id>/notas/<int:nid>/eliminar", methods=["POST"])
@login_required
def socio_nota_eliminar(id, nid):
    nota = AutoclubNota.query.get_or_404(nid)
    db.session.delete(nota)
    db.session.commit()
    flash("Nota eliminada.", "info")
    return redirect(url_for("autoclub.detalle", id=id) + "#notas")


@autoclub_bp.route("/patrocinadores/<int:id>/notas/nueva", methods=["POST"])
@login_required
def patrocinador_nota_nueva(id):
    p = Patrocinador.query.get_or_404(id)
    contenido = request.form.get("contenido", "").strip()
    if contenido:
        nota = AutoclubNota(
            entidad="patrocinador",
            entidad_id=p.id,
            entidad_nombre=p.nombre,
            contenido=contenido,
            usuario_nombre=current_user.username if hasattr(current_user, "username") else "",
        )
        db.session.add(nota)
        db.session.commit()
        flash("Nota añadida.", "success")
    return_to = request.form.get("return_to", "detalle")
    if return_to == "lista":
        return redirect(url_for("autoclub.patrocinadores_lista"))
    return redirect(url_for("autoclub.patrocinadores_detalle", id=id) + "#notas")


@autoclub_bp.route("/patrocinadores/<int:id>/notas/<int:nid>/eliminar", methods=["POST"])
@login_required
def patrocinador_nota_eliminar(id, nid):
    nota = AutoclubNota.query.get_or_404(nid)
    db.session.delete(nota)
    db.session.commit()
    flash("Nota eliminada.", "info")
    return redirect(url_for("autoclub.patrocinadores_detalle", id=id) + "#notas")


# ══════════════════════════════════════════════════════════════════════════════
#  INFORMACIÓN DEL AUTOCLUB
# ══════════════════════════════════════════════════════════════════════════════

@autoclub_bp.route("/informacion", methods=["GET", "POST"])
@login_required
def informacion():
    from app.models.autoclub_info import AutoclubInfo, AutoclubCredencial
    info = AutoclubInfo.get()
    if request.method == "POST":
        info.nombre_escuderia   = request.form.get("nombre_escuderia", "").strip()
        info.cif                = request.form.get("cif", "").strip()
        info.licencia_escuderia = request.form.get("licencia_escuderia", "").strip()
        info.email_contacto     = request.form.get("email_contacto", "").strip()
        info.telefono           = request.form.get("telefono", "").strip()
        info.direccion          = request.form.get("direccion", "").strip()
        info.notas              = request.form.get("notas", "").strip()
        db.session.commit()
        registrar_log("editar", "socio", 0, "Información del AutoClub actualizada")
        flash("Información del AutoClub guardada correctamente.", "success")
        return redirect(url_for("autoclub.informacion"))
    credenciales = AutoclubCredencial.query.order_by(AutoclubCredencial.nombre.asc()).all()
    return render_template("autoclub/informacion.html", info=info, credenciales=credenciales)


@autoclub_bp.route("/informacion/gastos/exportar")
@login_required
def informacion_gastos_exportar():
    import io
    from datetime import date as _date
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    def _parse(s):
        try:
            return _date.fromisoformat(s) if s else None
        except ValueError:
            return None

    fecha_desde = _parse(request.args.get("fecha_desde", ""))
    fecha_hasta = _parse(request.args.get("fecha_hasta", ""))

    q = GastoInscripcion.query
    if fecha_desde:
        q = q.filter(GastoInscripcion.fecha >= fecha_desde)
    if fecha_hasta:
        q = q.filter(GastoInscripcion.fecha <= fecha_hasta)
    gastos = q.order_by(GastoInscripcion.fecha.asc(), GastoInscripcion.creado_en.asc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Gastos"

    NEGRO         = "111111"
    VERDE         = "16A34A"
    GRIS_CABECERA = "F3F4F6"
    GRIS_BORDE    = "D1D5DB"

    def fill(hex_color):
        return PatternFill("solid", fgColor=hex_color)

    def borde_fino():
        lado = Side(style="thin", color=GRIS_BORDE)
        return Border(left=lado, right=lado, top=lado, bottom=lado)

    # ── Bloque de cabecera ───────────────────────────────────
    desde_str = fecha_desde.strftime("%d/%m/%Y") if fecha_desde else "—"
    hasta_str = fecha_hasta.strftime("%d/%m/%Y") if fecha_hasta else "—"
    meta = [
        ("Período", f"{desde_str}  →  {hasta_str}"),
        ("Total gastos", str(len(gastos))),
        ("Exportado", _date.today().strftime("%d/%m/%Y")),
    ]
    for r, (etiqueta, valor) in enumerate(meta, start=1):
        ws.cell(r, 1, etiqueta).font  = Font(bold=True, color=NEGRO, size=9)
        ws.cell(r, 1).fill            = fill(GRIS_CABECERA)
        ws.cell(r, 2, valor).font     = Font(color=NEGRO, size=9)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)

    ws.append([])  # fila vacía separadora

    # ── Cabecera de tabla (fila 5) ───────────────────────────
    HDR_ROW = 5
    headers = ["Prueba", "Concepto", "Categoría", "Fecha", "Importe (€)"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(HDR_ROW, c, h)
        cell.font      = Font(bold=True, color="FFFFFF", size=10)
        cell.fill      = fill(NEGRO)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = borde_fino()
    ws.row_dimensions[HDR_ROW].height = 22

    # ── Filas de datos ───────────────────────────────────────
    cat_colors = {
        "Inscripción": "DBEAFE", "Combustible": "FEF3C7",
        "Alojamiento": "F3E8FF", "Dietas":      "FCE7F3",
        "Transporte":  "E0F2FE", "Repuestos":   "FEE2E2",
        "Neumáticos":  "FFEDD5", "Equipación":  "ECFDF5",
        "Otros":       "F9FAFB",
    }
    total = 0.0
    for i, g in enumerate(gastos):
        row   = HDR_ROW + 1 + i
        zebra = "FFFFFF" if i % 2 == 0 else "F9FAFB"
        bg    = cat_colors.get(g.categoria or "Otros", "F9FAFB")

        prueba_nombre = g.inscripcion.nombre_prueba if g.inscripcion else ""
        prueba_cell = ws.cell(row, 1, prueba_nombre)
        prueba_cell.font   = Font(size=9, color=NEGRO)
        prueba_cell.fill   = fill(zebra)
        prueba_cell.border = borde_fino()

        concepto_cell = ws.cell(row, 2, g.concepto)
        concepto_cell.font   = Font(size=9)
        concepto_cell.fill   = fill(zebra)
        concepto_cell.border = borde_fino()

        cat_cell = ws.cell(row, 3, g.categoria or "")
        cat_cell.font      = Font(size=9, color=NEGRO)
        cat_cell.fill      = fill(bg)
        cat_cell.alignment = Alignment(horizontal="center")
        cat_cell.border    = borde_fino()

        fecha_cell = ws.cell(row, 4, g.fecha.strftime("%d/%m/%Y") if g.fecha else "")
        fecha_cell.font      = Font(size=9)
        fecha_cell.fill      = fill(zebra)
        fecha_cell.alignment = Alignment(horizontal="center")
        fecha_cell.border    = borde_fino()

        imp_cell = ws.cell(row, 5, g.importe)
        imp_cell.font          = Font(size=9)
        imp_cell.fill          = fill(zebra)
        imp_cell.number_format = '#,##0.00 "€"'
        imp_cell.alignment     = Alignment(horizontal="right")
        imp_cell.border        = borde_fino()
        total += g.importe

    # ── Fila de total ────────────────────────────────────────
    total_row = HDR_ROW + 1 + len(gastos) + 1
    ws.cell(total_row, 1, "TOTAL").font      = Font(bold=True, color="FFFFFF", size=10)
    ws.cell(total_row, 1).fill               = fill(VERDE)
    ws.cell(total_row, 1).alignment          = Alignment(horizontal="right")
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    imp_total = ws.cell(total_row, 5, total)
    imp_total.font          = Font(bold=True, color="FFFFFF", size=10)
    imp_total.fill          = fill(VERDE)
    imp_total.number_format = '#,##0.00 "€"'
    imp_total.alignment     = Alignment(horizontal="right")

    # ── Anchos de columna ────────────────────────────────────
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 16

    ws.freeze_panes = ws.cell(HDR_ROW + 1, 1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    nombre = f"gastos_autoclub_{_date.today().strftime('%Y%m%d')}.xlsx"
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
    )


@autoclub_bp.route("/informacion/credencial/nueva", methods=["POST"])
@login_required
def credencial_nueva():
    from app.models.autoclub_info import AutoclubCredencial
    cred = AutoclubCredencial(
        nombre   = request.form.get("nombre", "").strip(),
        url      = request.form.get("url", "").strip(),
        usuario  = request.form.get("usuario", "").strip(),
        password = request.form.get("password", "").strip(),
        notas    = request.form.get("notas", "").strip(),
    )
    if not cred.nombre:
        flash("El nombre del sitio es obligatorio.", "warning")
        return redirect(url_for("autoclub.informacion") + "#credenciales")
    db.session.add(cred)
    db.session.commit()
    registrar_log("crear", "socio", cred.id, f"Credencial añadida: {cred.nombre}")
    flash(f"Credencial «{cred.nombre}» añadida.", "success")
    return redirect(url_for("autoclub.informacion") + "#credenciales")


@autoclub_bp.route("/informacion/credencial/<int:id>/editar", methods=["POST"])
@login_required
def credencial_editar(id):
    from app.models.autoclub_info import AutoclubCredencial
    cred = AutoclubCredencial.query.get_or_404(id)
    cred.nombre   = request.form.get("nombre", "").strip()
    cred.url      = request.form.get("url", "").strip()
    cred.usuario  = request.form.get("usuario", "").strip()
    cred.password = request.form.get("password", "").strip()
    cred.notas    = request.form.get("notas", "").strip()
    db.session.commit()
    registrar_log("editar", "socio", id, f"Credencial editada: {cred.nombre}")
    flash(f"Credencial «{cred.nombre}» actualizada.", "success")
    return redirect(url_for("autoclub.informacion") + "#credenciales")


@autoclub_bp.route("/informacion/credencial/<int:id>/eliminar", methods=["POST"])
@login_required
def credencial_eliminar(id):
    from app.models.autoclub_info import AutoclubCredencial
    cred = AutoclubCredencial.query.get_or_404(id)
    nombre = cred.nombre
    db.session.delete(cred)
    db.session.commit()
    registrar_log("eliminar", "socio", id, f"Credencial eliminada: {nombre}")
    flash(f"Credencial «{nombre}» eliminada.", "info")
    return redirect(url_for("autoclub.informacion") + "#credenciales")


# ══════════════════════════════════════════════════════════════════════════════
#  PRENSA / CRÓNICA
# ══════════════════════════════════════════════════════════════════════════════

@autoclub_bp.route("/prensa")
@login_required
def prensa():
    from flask import current_app
    from app.models.medio_contacto import MedioContacto, CronicaEnviada
    from app.models.autoclub_info import AutoclubInfo
    medios = MedioContacto.query.order_by(MedioContacto.nombre.asc()).all()
    historial = CronicaEnviada.query.order_by(CronicaEnviada.enviado_en.desc()).limit(20).all()
    info = AutoclubInfo.get()
    mail_prensa_ok = bool(
        current_app.config.get("MAIL_PRENSA_USERNAME") and
        current_app.config.get("MAIL_PRENSA_PASSWORD")
    )
    return render_template("autoclub/prensa.html", medios=medios, historial=historial,
                           mail_prensa_ok=mail_prensa_ok, firma_guardada=info.firma_html or "")


@autoclub_bp.route("/prensa/test-smtp")
@login_required
def prensa_test_smtp():
    import smtplib, ssl as _ssl
    from flask import current_app
    user   = current_app.config.get("MAIL_PRENSA_USERNAME", "")
    passwd = current_app.config.get("MAIL_PRENSA_PASSWORD", "")
    server = current_app.config.get("MAIL_PRENSA_SERVER", "smtp.arsys.es")
    port   = current_app.config.get("MAIL_PRENSA_PORT", 465)
    pasos  = []
    try:
        ctx = _ssl.create_default_context()
        pasos.append(f"Conectando a {server}:{port}…")
        if port == 465:
            conn = smtplib.SMTP_SSL(server, port, context=ctx, timeout=10)
        else:
            conn = smtplib.SMTP(server, port, timeout=10)
        with conn as s:
            r = s.ehlo()
            pasos.append(f"EHLO → {r}")
            if port != 465:
                r = s.starttls(context=ctx)
                pasos.append(f"STARTTLS → {r}")
                r = s.ehlo()
                pasos.append(f"EHLO (post-TLS) → {r}")
            import base64 as _b64
            creds = _b64.b64encode(f"\x00{user}\x00{passwd}".encode()).decode()
            r = s.docmd("AUTH PLAIN", creds)
            pasos.append(f"AUTH PLAIN → {r}")
            if r[0] == 235:
                pasos.append("✅ Autenticación correcta — SMTP listo para enviar.")
            else:
                pasos.append(f"❌ Falló AUTH PLAIN, probando AUTH LOGIN…")
                r2 = s.login(user, passwd)
                pasos.append(f"AUTH LOGIN → {r2}")
                pasos.append("✅ Autenticación correcta vía LOGIN.")
    except smtplib.SMTPAuthenticationError as e:
        pasos.append(f"❌ SMTPAuthenticationError código {e.smtp_code}: {e.smtp_error}")
    except Exception as e:
        pasos.append(f"❌ {type(e).__name__}: {e}")
    return "<br>".join(str(p) for p in pasos)


@autoclub_bp.route("/prensa/medio/nuevo", methods=["POST"])
@login_required
def prensa_medio_nuevo():
    from app.models.medio_contacto import MedioContacto
    medio = MedioContacto(
        nombre = request.form.get("nombre", "").strip(),
        medio  = request.form.get("medio",  "").strip(),
        email  = request.form.get("email",  "").strip(),
        notas  = request.form.get("notas",  "").strip(),
    )
    if not medio.nombre or not medio.email:
        flash("Nombre y email son obligatorios.", "warning")
        return redirect(url_for("autoclub.prensa") + "#medios")
    db.session.add(medio)
    db.session.commit()
    registrar_log("crear", "socio", medio.id, f"Medio de prensa añadido: {medio.nombre}")
    flash(f"Contacto «{medio.nombre}» añadido.", "success")
    return redirect(url_for("autoclub.prensa") + "#medios")


@autoclub_bp.route("/prensa/medio/<int:id>/editar", methods=["POST"])
@login_required
def prensa_medio_editar(id):
    from app.models.medio_contacto import MedioContacto
    m = MedioContacto.query.get_or_404(id)
    m.nombre = request.form.get("nombre", "").strip()
    m.medio  = request.form.get("medio",  "").strip()
    m.email  = request.form.get("email",  "").strip()
    m.notas  = request.form.get("notas",  "").strip()
    db.session.commit()
    registrar_log("editar", "socio", id, f"Medio de prensa editado: {m.nombre}")
    flash(f"Contacto «{m.nombre}» actualizado.", "success")
    return redirect(url_for("autoclub.prensa") + "#medios")


@autoclub_bp.route("/prensa/medio/<int:id>/toggle", methods=["POST"])
@login_required
def prensa_medio_toggle(id):
    from app.models.medio_contacto import MedioContacto
    m = MedioContacto.query.get_or_404(id)
    m.activo = not m.activo
    db.session.commit()
    return redirect(url_for("autoclub.prensa") + "#medios")


@autoclub_bp.route("/prensa/medio/<int:id>/eliminar", methods=["POST"])
@login_required
def prensa_medio_eliminar(id):
    from app.models.medio_contacto import MedioContacto
    m = MedioContacto.query.get_or_404(id)
    nombre = m.nombre
    db.session.delete(m)
    db.session.commit()
    registrar_log("eliminar", "socio", id, f"Medio de prensa eliminado: {nombre}")
    flash(f"Contacto «{nombre}» eliminado.", "info")
    return redirect(url_for("autoclub.prensa") + "#medios")


@autoclub_bp.route("/prensa/enviar", methods=["POST"])
@login_required
def prensa_enviar():
    import json
    import re
    import time
    import random
    import smtplib
    import ssl
    import base64 as _b64
    from email.mime.multipart import MIMEMultipart
    from email.mime.text import MIMEText
    from email.mime.image import MIMEImage
    from email.utils import formatdate, make_msgid
    from flask import current_app, Response, stream_with_context
    from werkzeug.utils import secure_filename
    from app.models.medio_contacto import MedioContacto, CronicaEnviada
    from app.models.autoclub_info import AutoclubInfo

    def _ndjson(obj):
        return json.dumps(obj, ensure_ascii=False) + "\n"

    def _error_response(mensaje):
        return Response(_ndjson({"tipo": "error", "mensaje": mensaje}), mimetype="application/x-ndjson")

    asunto      = request.form.get("asunto", "").strip()
    cuerpo_html = request.form.get("cuerpo_html", "").strip()
    firma_html  = request.form.get("firma_html", "").strip()
    ids_raw     = request.form.getlist("destinatarios")

    # Persistir la firma para la próxima vez
    if firma_html:
        info = AutoclubInfo.get()
        info.firma_html = firma_html
        db.session.commit()

    if not asunto or not cuerpo_html:
        return _error_response("El asunto y el cuerpo son obligatorios.")

    if not ids_raw:
        return _error_response("Selecciona al menos un destinatario.")

    medios = MedioContacto.query.filter(MedioContacto.id.in_([int(i) for i in ids_raw])).all()
    if not medios:
        return _error_response("No se encontraron los destinatarios seleccionados.")

    # Fotos adjuntas: se leen una sola vez en memoria (los streams de Flask
    # solo se pueden consumir una vez) para reutilizarlas en cada email.
    fotos_adjuntas = []
    for f in request.files.getlist("fotos"):
        if not f or not f.filename:
            continue
        ext = f.filename.rsplit(".", 1)[-1].lower() if "." in f.filename else ""
        if ext not in FOTO_PRENSA_EXTENSIONS:
            return _error_response(f"Formato de imagen no admitido: {f.filename}. Usa JPG, PNG, WEBP o GIF.")
        fotos_adjuntas.append({
            "filename": secure_filename(f.filename) or f"foto.{ext}",
            "data": f.read(),
            "subtype": FOTO_PRENSA_MIME[ext],
        })

    prensa_user   = current_app.config.get("MAIL_PRENSA_USERNAME", "")
    prensa_pass   = current_app.config.get("MAIL_PRENSA_PASSWORD", "")
    prensa_sender = current_app.config.get("MAIL_PRENSA_SENDER", "") or prensa_user
    smtp_server   = current_app.config.get("MAIL_PRENSA_SERVER", "smtp.arsys.es")
    smtp_port     = current_app.config.get("MAIL_PRENSA_PORT", 465)
    sender_domain = prensa_user.split("@")[-1] if "@" in prensa_user else "mail.local"

    if not prensa_user or not prensa_pass:
        return _error_response("Configura MAIL_PRENSA_USERNAME y MAIL_PRENSA_PASSWORD en el archivo .env antes de enviar.")

    # Se capturan como valores planos antes del generador: current_user es un
    # proxy ligado al contexto de la petición y app_context (necesario para
    # current_app/db dentro del generador) se gestiona con stream_with_context.
    enviado_por = current_user.username or current_user.nombre

    SEPARADOR_FIRMA = (
        '<hr style="border:none;border-top:1px solid #e5e7eb;margin:32px 0 16px">'
    )

    def _html_a_texto(html):
        """Versión texto plano mínima para la parte alternativa."""
        txt = re.sub(r'<br\s*/?>', '\n', html, flags=re.I)
        txt = re.sub(r'<p[^>]*>', '\n', txt, flags=re.I)
        txt = re.sub(r'</p>', '\n', txt, flags=re.I)
        txt = re.sub(r'<hr[^>]*>', '\n---\n', txt, flags=re.I)
        txt = re.sub(r'<[^>]+>', '', txt)
        txt = re.sub(r'\n{3,}', '\n\n', txt)
        return txt.strip()

    def _personalizar(texto, medio_obj):
        return (texto
                .replace("{{nombre}}", medio_obj.nombre)
                .replace("{{medio}}",  medio_obj.medio or medio_obj.nombre))

    def _build_mime(medio_obj):
        cuerpo_final = _personalizar(cuerpo_html, medio_obj)
        if firma_html:
            cuerpo_final += SEPARADOR_FIRMA + _personalizar(firma_html, medio_obj)

        # Con fotos adjuntas necesitamos un contenedor "mixed" por fuera y el
        # texto/html "alternative" anidado dentro; sin fotos, se mantiene el
        # alternative simple de siempre (mensaje más ligero).
        alt = MIMEMultipart("alternative")
        # Los filtros antispam valoran que exista la versión texto además del HTML
        alt.attach(MIMEText(_html_a_texto(cuerpo_final), "plain", "utf-8"))
        alt.attach(MIMEText(cuerpo_final, "html", "utf-8"))

        outer = MIMEMultipart("mixed") if fotos_adjuntas else alt
        if fotos_adjuntas:
            outer.attach(alt)
            for foto in fotos_adjuntas:
                img = MIMEImage(foto["data"], _subtype=foto["subtype"])
                img.add_header("Content-Disposition", "attachment", filename=foto["filename"])
                outer.attach(img)

        outer["Subject"]      = asunto
        outer["From"]         = prensa_sender
        outer["To"]           = medio_obj.email
        outer["Date"]         = formatdate(localtime=True)
        outer["Message-ID"]   = make_msgid(domain=sender_domain)
        # List-Unsubscribe reduce la probabilidad de marcado como spam manual
        outer["List-Unsubscribe"] = f"<mailto:{prensa_user}?subject=Unsubscribe>"
        outer["List-Unsubscribe-Post"] = "List-Unsubscribe=One-Click"
        return outer

    def _conectar_smtp():
        context = ssl.create_default_context()
        if smtp_port == 465:
            conn = smtplib.SMTP_SSL(smtp_server, smtp_port, context=context, timeout=20)
        else:
            conn = smtplib.SMTP(smtp_server, smtp_port, timeout=20)
            conn.ehlo()
            conn.starttls(context=context)
        conn.ehlo()
        credentials = _b64.b64encode(
            f"\x00{prensa_user}\x00{prensa_pass}".encode("utf-8")
        ).decode()
        code, msg = conn.docmd("AUTH PLAIN", credentials)
        if code != 235:
            try:
                conn.close()
            except Exception:
                pass
            raise smtplib.SMTPAuthenticationError(code, msg)
        return conn

    def _generar():
        ok = 0
        err = 0
        dest_log = []

        try:
            server = _conectar_smtp()
        except smtplib.SMTPAuthenticationError as e:
            yield _ndjson({"tipo": "error", "mensaje":
                f"Error de autenticación SMTP (código {e.smtp_code}): {e.smtp_error.decode(errors='replace')} "
                f"— Revisa usuario y contraseña en el .env."})
            return
        except Exception as e:
            yield _ndjson({"tipo": "error", "mensaje": f"Error de conexión SMTP: {type(e).__name__}: {e}"})
            return

        try:
            for i, m in enumerate(medios):
                yield _ndjson({"tipo": "progreso", "actual": i + 1, "total": len(medios),
                                "nombre": m.nombre, "email": m.email})
                # El servidor de correo puede cortar la conexión a mitad de la
                # tanda (límite de envíos, timeout de inactividad, etc.). Si eso
                # ocurre, antes se seguía usando la misma conexión rota para el
                # resto de destinatarios y un único corte se propagaba en cascada
                # a todos los que quedaban. Ahora se reconecta y se reintenta una
                # vez antes de dar por perdido a ese destinatario.
                for intento in (1, 2):
                    try:
                        server.send_message(_build_mime(m))
                        ok += 1
                        dest_log.append({"nombre": m.nombre, "medio": m.medio, "email": m.email, "ok": True})
                        break
                    except (smtplib.SMTPServerDisconnected, smtplib.SMTPConnectError, OSError) as e:
                        if intento == 2:
                            err += 1
                            dest_log.append({"nombre": m.nombre, "medio": m.medio, "email": m.email,
                                              "ok": False, "error": f"Conexión perdida: {e}"})
                            break
                        try:
                            server.close()
                        except Exception:
                            pass
                        try:
                            server = _conectar_smtp()
                        except Exception as e2:
                            err += 1
                            dest_log.append({"nombre": m.nombre, "medio": m.medio, "email": m.email,
                                              "ok": False, "error": f"No se pudo reconectar: {e2}"})
                            break
                    except Exception as e:
                        err += 1
                        dest_log.append({"nombre": m.nombre, "medio": m.medio, "email": m.email, "ok": False, "error": str(e)})
                        break
                yield _ndjson({"tipo": "resultado", **dest_log[-1]})
                # Pausa aleatoria entre envíos: evita que el servidor saliente
                # acumule ráfagas que los receptores interpreten como spam
                if i < len(medios) - 1:
                    time.sleep(random.uniform(3, 7))
        finally:
            try:
                server.quit()
            except Exception:
                pass

        cronica = CronicaEnviada(
            asunto         = asunto,
            cuerpo_html    = cuerpo_html,
            destinatarios  = json.dumps(dest_log, ensure_ascii=False),
            total_enviados = ok,
            total_errores  = err,
            enviado_por    = enviado_por,
        )
        db.session.add(cronica)
        db.session.commit()
        registrar_log("crear", "inscripcion", cronica.id,
                      f"Crónica enviada: '{asunto}' → {ok} ok, {err} errores")

        if err == 0:
            mensaje = f"Crónica enviada correctamente a {ok} medio{'s' if ok != 1 else ''}."
        else:
            errores_detalle = "; ".join(
                f"{d['email']}: {d.get('error','?')}"
                for d in dest_log if not d.get('ok')
            )
            mensaje = f"Enviada a {ok} medios. {err} error{'es' if err != 1 else ''}: {errores_detalle}"

        yield _ndjson({"tipo": "fin", "ok": ok, "err": err, "mensaje": mensaje})

    return Response(stream_with_context(_generar()), mimetype="application/x-ndjson")


# ══════════════════════════════════════════════════════════════════════════════
#  NOTIFICACIONES INTERNAS
# ══════════════════════════════════════════════════════════════════════════════

@autoclub_bp.route("/notificaciones/<int:id>/leer", methods=["POST"])
@login_required
def notificacion_leer(id):
    from app.models.notificacion import Notificacion
    n = Notificacion.query.get_or_404(id)
    n.leida = True
    db.session.commit()
    return redirect(n.url or url_for("autoclub.inscripciones_lista"))


@autoclub_bp.route("/notificaciones/leer-todas", methods=["POST"])
@login_required
def notificaciones_leer_todas():
    from app.models.notificacion import Notificacion
    Notificacion.query.filter_by(leida=False).update({"leida": True})
    db.session.commit()
    next_url = request.referrer or url_for("autoclub.seguimiento")
    return redirect(next_url)


# ══════════════════════════════════════════════════════════════════════════════
#  SEGUIMIENTO / LOGS
# ══════════════════════════════════════════════════════════════════════════════

@autoclub_bp.route("/calendario")
@login_required
def calendario():
    from datetime import date
    from app.models.competicion_evento import CompeticionEvento
    hoy = date.today()
    proximas_insc = (Inscripcion.query
                     .filter(Inscripcion.fecha_prueba >= hoy)
                     .order_by(Inscripcion.fecha_prueba.asc())
                     .limit(8).all())
    proximos_ev = (CompeticionEvento.query
                   .filter(CompeticionEvento.fecha >= hoy)
                   .order_by(CompeticionEvento.fecha.asc())
                   .limit(8).all())
    return render_template("autoclub/calendario.html",
                           inscripciones_proximas=proximas_insc,
                           proximos_eventos=proximos_ev)


@autoclub_bp.route("/calendario/eventos")
@login_required
def calendario_eventos():
    from app.models.competicion_evento import CompeticionEvento
    COLORES = {
        "pendiente": "#6b7280",
        "enviada":   "#1d4ed8",
        "aceptada":  "#16a34a",
        "rechazada": "#dc2626",
        "parcial":   "#d97706",
    }
    eventos = []

    for insc in Inscripcion.query.filter(Inscripcion.fecha_prueba.isnot(None)).all():
        color = COLORES.get(insc.estado_grupo, "#6b7280")
        evento_json = {
            "id":    f"insc-{insc.id}",
            "title": insc.nombre_prueba,
            "start": insc.fecha_prueba.isoformat(),
            "allDay": True,
            "color": color,
            "extendedProps": {
                "tipo":       "inscripcion",
                "campeonato": insc.campeonato or "",
                "estado":     insc.estado_grupo,
                "piloto":     insc.piloto.nombre_completo if insc.piloto else "",
                "vehiculo":   insc.vehiculo.nombre_display if insc.vehiculo else "",
                "num_pilotos": insc.num_pilotos,
            },
            "url": url_for("autoclub.inscripciones_detalle", id=insc.id),
        }
        # FullCalendar trata "end" como exclusivo, así que para que el último
        # día de la prueba se pinte también hay que sumarle un día.
        if insc.fecha_fin and insc.fecha_fin > insc.fecha_prueba:
            evento_json["end"] = (insc.fecha_fin + timedelta(days=1)).isoformat()
        eventos.append(evento_json)

    for ev in CompeticionEvento.query.all():
        props = {
            "tipo":        "evento",
            "evento_id":   ev.id,
            "descripcion": ev.descripcion or "",
            "fecha_plazo": ev.fecha_plazo.isoformat() if ev.fecha_plazo else "",
        }
        eventos.append({
            "id":         f"ev-{ev.id}",
            "title":      ev.titulo,
            "start":      ev.fecha.isoformat(),
            "color":      "#f59e0b",
            "textColor":  "#111",
            "extendedProps": props,
            "url": "",
        })
        # Añadir marcador visual del plazo en el calendario
        if ev.fecha_plazo:
            eventos.append({
                "id":         f"ev-plazo-{ev.id}",
                "title":      f"⏰ Plazo: {ev.titulo}",
                "start":      ev.fecha_plazo.isoformat(),
                "color":      "#dc2626",
                "textColor":  "#fff",
                "display":    "list-item",
                "extendedProps": {"tipo": "plazo_evento", "evento_id": ev.id},
                "url": "",
            })

    return jsonify(eventos)


@autoclub_bp.route("/calendario/evento/nuevo", methods=["POST"])
@login_required
def calendario_evento_nuevo():
    from datetime import date as _date
    from app.models.competicion_evento import CompeticionEvento
    data = request.get_json(silent=True) or {}
    titulo       = (data.get("titulo")       or "").strip()
    descripcion  = (data.get("descripcion")  or "").strip()
    fecha_str    = (data.get("fecha")        or "").strip()
    plazo_str    = (data.get("fecha_plazo")  or "").strip()
    if not titulo or not fecha_str:
        return jsonify({"ok": False, "error": "Faltan datos"}), 400
    try:
        fecha = _date.fromisoformat(fecha_str)
    except ValueError:
        return jsonify({"ok": False, "error": "Fecha inválida"}), 400
    fecha_plazo = None
    if plazo_str:
        try:
            fecha_plazo = _date.fromisoformat(plazo_str)
        except ValueError:
            pass
    ev = CompeticionEvento(titulo=titulo, descripcion=descripcion,
                           fecha=fecha, fecha_plazo=fecha_plazo)
    db.session.add(ev)
    db.session.commit()
    registrar_log("crear", "inscripcion", ev.id, f"Evento de competición anotado: {ev.titulo}")
    return jsonify({"ok": True, "id": ev.id})


@autoclub_bp.route("/calendario/evento/<int:id>/editar", methods=["POST"])
@login_required
def calendario_evento_editar(id):
    from datetime import date as _date
    from app.models.competicion_evento import CompeticionEvento
    ev   = CompeticionEvento.query.get_or_404(id)
    data = request.get_json(silent=True) or {}
    titulo    = (data.get("titulo")      or "").strip()
    fecha_str = (data.get("fecha")       or "").strip()
    plazo_str = (data.get("fecha_plazo") or "").strip()
    desc      = (data.get("descripcion") or "").strip()
    if not titulo or not fecha_str:
        return jsonify({"ok": False, "error": "Faltan datos"}), 400
    try:
        ev.fecha = _date.fromisoformat(fecha_str)
    except ValueError:
        return jsonify({"ok": False, "error": "Fecha inválida"}), 400
    ev.titulo      = titulo
    ev.descripcion = desc
    ev.fecha_plazo = None
    if plazo_str:
        try:
            ev.fecha_plazo = _date.fromisoformat(plazo_str)
        except ValueError:
            pass
    db.session.commit()
    registrar_log("editar", "inscripcion", ev.id, f"Evento de competición editado: {ev.titulo}")
    return jsonify({"ok": True})


@autoclub_bp.route("/calendario/evento/<int:id>/eliminar", methods=["POST"])
@login_required
def calendario_evento_eliminar(id):
    from app.models.competicion_evento import CompeticionEvento
    ev = CompeticionEvento.query.get_or_404(id)
    nombre = ev.titulo
    db.session.delete(ev)
    db.session.commit()
    registrar_log("eliminar", "inscripcion", id, f"Evento de competición eliminado: {nombre}")
    return jsonify({"ok": True})


@autoclub_bp.route("/seguimiento")
@login_required
def seguimiento():
    from app.models.log import Log
    from app.models.usuario import Usuario
    from datetime import datetime

    page      = request.args.get("page", 1, type=int)
    usuario   = request.args.get("usuario", "").strip()
    accion    = request.args.get("accion", "").strip()
    fecha_ini = request.args.get("fecha_ini", "").strip()
    fecha_fin = request.args.get("fecha_fin", "").strip()

    ENTIDADES_AUTOCLUB = ["socio", "patrocinador", "piloto", "vehiculo", "inscripcion", "staff"]
    entidad   = request.args.get("entidad", "").strip()

    base_filter = db.or_(
        Log.accion.in_(["login", "logout", "login_fallido"]),
        Log.entidad.in_(ENTIDADES_AUTOCLUB),
    )
    query = Log.query.filter(base_filter)

    if entidad:
        query = query.filter(Log.entidad == entidad)
    if usuario:
        query = query.filter(Log.usuario_nombre.ilike(f"%{usuario}%"))
    if accion:
        query = query.filter(Log.accion == accion)
    if fecha_ini:
        try:
            query = query.filter(Log.fecha >= datetime.strptime(fecha_ini, "%Y-%m-%d"))
        except ValueError:
            pass
    if fecha_fin:
        from datetime import timedelta
        try:
            query = query.filter(
                Log.fecha < datetime.strptime(fecha_fin, "%Y-%m-%d") + timedelta(days=1)
            )
        except ValueError:
            pass

    logs = query.order_by(Log.fecha.desc()).paginate(page=page, per_page=40, error_out=False)

    acciones_disponibles = ["login", "logout", "login_fallido",
                            "crear", "editar", "eliminar", "cambiar_estado"]
    usuarios_distintos = sorted({
        r[0] for r in db.session.query(Log.usuario_nombre).filter(base_filter).distinct().all()
    })

    return render_template(
        "autoclub/seguimiento.html",
        logs=logs,
        acciones=acciones_disponibles,
        usuarios=usuarios_distintos,
        entidades=ENTIDADES_AUTOCLUB,
        filtro_usuario=usuario,
        filtro_accion=accion,
        filtro_entidad=entidad,
        filtro_fecha_ini=fecha_ini,
        filtro_fecha_fin=fecha_fin,
    )


# ══════════════════════════════════════════════════════
# SELECCIÓN FEMENINA — panel admin de preinscripciones
# (la landing pública + el alta viven en app/routes/seleccion_femenina.py,
# sin login; esto es solo la vista de gestión para el staff)
# ══════════════════════════════════════════════════════

@autoclub_bp.route("/seleccion-femenina")
@login_required
def seleccion_femenina_lista():
    estado = request.args.get("estado", "")
    q = request.args.get("q", "").strip()
    query = PreinscripcionCarcross.query
    if estado:
        query = query.filter_by(estado=estado)
    if q:
        like = f"%{q}%"
        query = query.filter(db.or_(
            PreinscripcionCarcross.nombre_completo.ilike(like),
            PreinscripcionCarcross.email.ilike(like),
            PreinscripcionCarcross.localidad.ilike(like),
        ))
    preinscripciones = query.order_by(PreinscripcionCarcross.creado_en.desc()).all()
    return render_template(
        "autoclub/seleccion_femenina/lista.html",
        preinscripciones=preinscripciones, estado=estado, q=q,
        estados=ESTADOS_PREINSCRIPCION,
    )


@autoclub_bp.route("/seleccion-femenina/<int:id>/estado", methods=["POST"])
@login_required
def seleccion_femenina_estado(id):
    preinscripcion = PreinscripcionCarcross.query.get_or_404(id)
    nuevo_estado = request.form.get("estado")
    if nuevo_estado in ESTADOS_PREINSCRIPCION:
        disparar_email_inscripcion = (
            nuevo_estado == "apta_inscripcion" and preinscripcion.estado != "apta_inscripcion"
        )
        preinscripcion.estado = nuevo_estado
        db.session.commit()
        registrar_log("cambiar_estado", "preinscripcion_carcross", id,
                      f"Estado -> {nuevo_estado}: {preinscripcion.nombre_completo}")

        if disparar_email_inscripcion:
            from app.services.inscripcion_autorizacion_service import obtener_o_crear_inscripcion, enviar_email_inscripcion
            inscripcion = obtener_o_crear_inscripcion(preinscripcion)
            if enviar_email_inscripcion(preinscripcion, inscripcion):
                flash("Estado actualizado y email de inscripción/autorización enviado.", "success")
            else:
                flash("Estado actualizado, pero falló el envío del email de inscripción (revisa el historial de mails).", "warning")
        else:
            flash("Estado actualizado.", "success")
    return redirect(request.referrer or url_for("autoclub.seleccion_femenina_lista"))


@autoclub_bp.route("/seleccion-femenina/<int:id>/notas", methods=["POST"])
@login_required
def seleccion_femenina_notas(id):
    preinscripcion = PreinscripcionCarcross.query.get_or_404(id)
    preinscripcion.notas_internas = request.form.get("notas_internas", "").strip()
    db.session.commit()
    flash("Notas guardadas.", "success")
    return redirect(request.referrer or url_for("autoclub.seleccion_femenina_lista"))


@autoclub_bp.route("/seleccion-femenina/exportar")
@login_required
def seleccion_femenina_exportar():
    preinscripciones = PreinscripcionCarcross.query.order_by(PreinscripcionCarcross.creado_en.desc()).all()
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(["ID", "Nombre", "Edad", "Fecha_nacimiento", "Localidad", "Email",
                     "Telefono", "Experiencia_previa", "Motivacion", "Estado", "Recibida"])
    for p in preinscripciones:
        writer.writerow([
            p.id, p.nombre_completo, p.edad,
            p.fecha_nacimiento.strftime("%d/%m/%Y") if p.fecha_nacimiento else "",
            p.localidad, p.email, p.telefono, p.experiencia_previa_label,
            p.motivacion, p.estado, p.creado_en.strftime("%d/%m/%Y %H:%M") if p.creado_en else "",
        ])
    resp = Response(output.getvalue(), mimetype="text/csv; charset=utf-8")
    resp.headers["Content-Disposition"] = "attachment; filename=preinscripciones_seleccion_femenina.csv"
    return resp


def _leer_form_preinscripcion(form, errores):
    nombre = form.get("nombre_completo", "").strip()
    fnac_str = form.get("fecha_nacimiento", "").strip()
    localidad = form.get("localidad", "").strip()
    email = form.get("email", "").strip()
    telefono = form.get("telefono", "").strip()
    experiencia = form.get("experiencia_previa", "ninguna")
    motivacion = form.get("motivacion", "").strip()
    notas_internas = form.get("notas_internas", "").strip()
    estado = form.get("estado", "pendiente")
    altura_str = form.get("altura_cm", "").strip()
    peso_str = form.get("peso_kg", "").strip()
    talla_camiseta = form.get("talla_camiseta", "").strip()
    talla_zapatillas = form.get("talla_zapatillas", "").strip()

    if not nombre:
        errores["nombre_completo"] = "Indica el nombre y apellidos."
    fecha_nacimiento = None
    if not fnac_str:
        errores["fecha_nacimiento"] = "Indica la fecha de nacimiento."
    else:
        try:
            fecha_nacimiento = datetime.strptime(fnac_str, "%Y-%m-%d").date()
        except ValueError:
            errores["fecha_nacimiento"] = "Fecha no válida."
    if not email:
        errores["email"] = "Indica un email."
    if not telefono:
        errores["telefono"] = "Indica un teléfono."
    if experiencia not in dict(EXPERIENCIA_PREVIA_OPCIONES):
        experiencia = "ninguna"
    if estado not in ESTADOS_PREINSCRIPCION:
        estado = "pendiente"

    altura_cm = None
    if altura_str:
        try:
            altura_cm = int(altura_str)
        except ValueError:
            errores["altura_cm"] = "Indica la altura en centímetros (solo número)."

    peso_kg = None
    if peso_str:
        try:
            peso_kg = int(peso_str)
        except ValueError:
            errores["peso_kg"] = "Indica el peso en kilos (solo número)."

    return {
        "nombre_completo": nombre, "fecha_nacimiento": fecha_nacimiento, "localidad": localidad,
        "email": email, "telefono": telefono, "experiencia_previa": experiencia,
        "motivacion": motivacion, "notas_internas": notas_internas, "estado": estado,
        "altura_cm": altura_cm, "peso_kg": peso_kg,
        "talla_camiseta": talla_camiseta, "talla_zapatillas": talla_zapatillas,
    }


def _datos_form_repoblar(form):
    """Dict de strings tal cual se escribieron en el formulario, para
    volver a rellenarlo si falla la validación — nunca se pasa
    request.form directamente a la plantilla (es un MultiDict, no un dict
    plano, y fecha_nacimiento necesita quedar siempre como string ISO para
    el <input type=date>, igual que al precargar una preinscripción ya
    guardada en base de datos)."""
    return {
        "nombre_completo": form.get("nombre_completo", ""),
        "fecha_nacimiento": form.get("fecha_nacimiento", ""),
        "localidad": form.get("localidad", ""),
        "email": form.get("email", ""),
        "telefono": form.get("telefono", ""),
        "experiencia_previa": form.get("experiencia_previa", "ninguna"),
        "motivacion": form.get("motivacion", ""),
        "notas_internas": form.get("notas_internas", ""),
        "estado": form.get("estado", "pendiente"),
        "altura_cm": form.get("altura_cm", ""),
        "peso_kg": form.get("peso_kg", ""),
        "talla_camiseta": form.get("talla_camiseta", ""),
        "talla_zapatillas": form.get("talla_zapatillas", ""),
    }


def _datos_desde_preinscripcion(p):
    return {
        "nombre_completo": p.nombre_completo,
        "fecha_nacimiento": p.fecha_nacimiento.isoformat() if p.fecha_nacimiento else "",
        "localidad": p.localidad or "",
        "email": p.email,
        "telefono": p.telefono,
        "experiencia_previa": p.experiencia_previa,
        "motivacion": p.motivacion or "",
        "notas_internas": p.notas_internas or "",
        "estado": p.estado,
        "altura_cm": p.altura_cm if p.altura_cm is not None else "",
        "peso_kg": p.peso_kg if p.peso_kg is not None else "",
        "talla_camiseta": p.talla_camiseta or "",
        "talla_zapatillas": p.talla_zapatillas or "",
    }


@autoclub_bp.route("/seleccion-femenina/nueva", methods=["GET", "POST"])
@login_required
def seleccion_femenina_nueva():
    if request.method == "POST":
        errores = {}
        datos = _leer_form_preinscripcion(request.form, errores)
        if errores:
            flash("Revisa los campos marcados.", "danger")
            return render_template(
                "autoclub/seleccion_femenina/form.html", preinscripcion=None,
                datos=_datos_form_repoblar(request.form), errores=errores,
                experiencia_opciones=EXPERIENCIA_PREVIA_OPCIONES, estados=ESTADOS_PREINSCRIPCION,
            ), 400

        preinscripcion = PreinscripcionCarcross(acepta_privacidad=True, ip="alta-manual-crm", **datos)
        db.session.add(preinscripcion)
        db.session.commit()
        registrar_log("crear", "preinscripcion_carcross", preinscripcion.id,
                      f"Alta manual: {preinscripcion.nombre_completo}")
        flash("Preinscripción creada.", "success")
        return redirect(url_for("autoclub.seleccion_femenina_lista"))

    return render_template(
        "autoclub/seleccion_femenina/form.html", preinscripcion=None,
        datos=None, errores=None,
        experiencia_opciones=EXPERIENCIA_PREVIA_OPCIONES, estados=ESTADOS_PREINSCRIPCION,
    )


@autoclub_bp.route("/seleccion-femenina/<int:id>/editar", methods=["GET", "POST"])
@login_required
def seleccion_femenina_editar(id):
    preinscripcion = PreinscripcionCarcross.query.get_or_404(id)
    if request.method == "POST":
        errores = {}
        datos = _leer_form_preinscripcion(request.form, errores)
        if errores:
            flash("Revisa los campos marcados.", "danger")
            return render_template(
                "autoclub/seleccion_femenina/form.html", preinscripcion=preinscripcion,
                datos=_datos_form_repoblar(request.form), errores=errores,
                experiencia_opciones=EXPERIENCIA_PREVIA_OPCIONES, estados=ESTADOS_PREINSCRIPCION,
                mailings=preinscripcion.mailings.all(),
            ), 400

        for campo, valor in datos.items():
            setattr(preinscripcion, campo, valor)
        db.session.commit()
        registrar_log("editar", "preinscripcion_carcross", preinscripcion.id,
                      f"Preinscripción editada: {preinscripcion.nombre_completo}")
        flash("Preinscripción actualizada.", "success")
        return redirect(url_for("autoclub.seleccion_femenina_lista"))

    return render_template(
        "autoclub/seleccion_femenina/form.html", preinscripcion=preinscripcion,
        datos=_datos_desde_preinscripcion(preinscripcion), errores=None,
        experiencia_opciones=EXPERIENCIA_PREVIA_OPCIONES, estados=ESTADOS_PREINSCRIPCION,
        mailings=preinscripcion.mailings.all(),
    )


@autoclub_bp.route("/seleccion-femenina/<int:id>/inscripcion")
@login_required
def seleccion_femenina_inscripcion_detalle(id):
    preinscripcion = PreinscripcionCarcross.query.get_or_404(id)
    if not preinscripcion.inscripcion_autorizacion:
        abort(404)
    return render_template(
        "autoclub/seleccion_femenina/inscripcion_detalle.html",
        preinscripcion=preinscripcion, insc=preinscripcion.inscripcion_autorizacion,
    )


@autoclub_bp.route("/seleccion-femenina/<int:id>/inscripcion/pdf")
@login_required
def seleccion_femenina_inscripcion_pdf(id):
    from app.services.inscripcion_pdf import generar_pdf_inscripcion
    preinscripcion = PreinscripcionCarcross.query.get_or_404(id)
    if not preinscripcion.inscripcion_autorizacion or not preinscripcion.inscripcion_autorizacion.completo:
        abort(404)
    pdf_bytes = generar_pdf_inscripcion(preinscripcion, preinscripcion.inscripcion_autorizacion)
    return Response(pdf_bytes, mimetype="application/pdf", headers={
        "Content-Disposition": f"inline; filename=inscripcion_autorizacion_{preinscripcion.id}.pdf",
    })


@autoclub_bp.route("/seleccion-femenina/<int:id>/eliminar", methods=["POST"])
@login_required
def seleccion_femenina_eliminar(id):
    preinscripcion = PreinscripcionCarcross.query.get_or_404(id)
    nombre = preinscripcion.nombre_completo
    db.session.delete(preinscripcion)
    db.session.commit()
    registrar_log("eliminar", "preinscripcion_carcross", id, f"Preinscripción eliminada: {nombre}")
    flash(f"Preinscripción de «{nombre}» eliminada.", "info")
    return redirect(url_for("autoclub.seleccion_femenina_lista"))


@autoclub_bp.route("/seleccion-femenina/mailing", methods=["GET", "POST"])
@login_required
def seleccion_femenina_mailing():
    from app.services.seleccion_femenina_mail import destinatarias_mailing, _html_mailing

    if request.method == "POST" and request.form.get("accion") == "previsualizar":
        asunto = request.form.get("asunto", "").strip()
        cuerpo = request.form.get("cuerpo", "").strip()
        estados_sel = request.form.getlist("estados")
        if not asunto or not cuerpo:
            flash("Indica un asunto y un cuerpo para el email.", "danger")
            return redirect(url_for("autoclub.seleccion_femenina_mailing"))

        destinatarias = destinatarias_mailing(estados_sel or None)
        preview_html = _html_mailing("Nombre de ejemplo", cuerpo.replace("\n", "<br>"))
        return render_template(
            "autoclub/seleccion_femenina/mailing.html",
            estados=ESTADOS_PREINSCRIPCION, estados_sel=estados_sel,
            asunto=asunto, cuerpo=cuerpo,
            preview_html=preview_html, destinatarias=destinatarias,
            resultado=None,
        )

    return render_template(
        "autoclub/seleccion_femenina/mailing.html",
        estados=ESTADOS_PREINSCRIPCION, estados_sel=[],
        asunto="", cuerpo="", preview_html=None, destinatarias=None,
        resultado=None,
    )


@autoclub_bp.route("/seleccion-femenina/mailing/enviar", methods=["POST"])
@login_required
def seleccion_femenina_mailing_enviar():
    from app.services.seleccion_femenina_mail import enviar_mailing_masivo

    asunto = request.form.get("asunto", "").strip()
    cuerpo = request.form.get("cuerpo", "").strip()
    estados_sel = request.form.getlist("estados")
    if not asunto or not cuerpo:
        flash("Indica un asunto y un cuerpo para el email.", "danger")
        return redirect(url_for("autoclub.seleccion_femenina_mailing"))

    resultado = enviar_mailing_masivo(asunto, cuerpo, estados_sel or None)
    registrar_log("crear", "preinscripcion_carcross", None, (
        f"Mailing enviado: {len(resultado['enviados'])} correctos, "
        f"{len(resultado['fallidos'])} fallidos — asunto: {asunto}"
    ))
    if resultado["fallidos"]:
        flash(
            f"{len(resultado['enviados'])} emails enviados. {len(resultado['fallidos'])} fallaron.",
            "warning",
        )
    else:
        flash(f"{len(resultado['enviados'])} emails enviados correctamente.", "success")

    return render_template(
        "autoclub/seleccion_femenina/mailing.html",
        estados=ESTADOS_PREINSCRIPCION, estados_sel=[],
        asunto="", cuerpo="", preview_html=None, destinatarias=None,
        resultado=resultado,
    )


# ══════════════════════════════════════════════════════════════════════════
# FIRMABLES — 3 documentos (aptitud médica, asunción de riesgo, disclaimer
# de conducción) rellenados presencialmente el día de la entrevista, en un
# wizard secuencial. Se completan en el propio dispositivo del staff, que
# se entrega a los progenitores/tutores presentes.
# ══════════════════════════════════════════════════════════════════════════

def _doc_vigente(preinscripcion_id, tipo):
    return (DocumentoFirmado.query
            .filter_by(preinscripcion_id=preinscripcion_id, tipo=tipo, vigente=True)
            .order_by(DocumentoFirmado.version.desc()).first())


@autoclub_bp.route("/seleccion-femenina/<int:id>/firmables")
@login_required
def seleccion_femenina_firmables(id):
    from app.services.firmables_service import enlace_firmables, obtener_o_crear_sesion_pendiente
    preinscripcion = PreinscripcionCarcross.query.get_or_404(id)
    estado_docs = [
        {"tipo": tipo, "label": label, "doc": _doc_vigente(id, tipo)}
        for tipo, label in TIPOS_FIRMABLE
    ]
    sesion_activa = SesionFirmables.query.filter_by(preinscripcion_id=id, estado="en_progreso").first()
    if sesion_activa:
        # Sesiones en_progreso creadas antes de que existiera el enlace por
        # token (o cuyo enlace ya caducó) se curan aquí regenerándolo, para
        # no dejar la pantalla mostrando "None" sin forma de arreglarlo.
        sesion_activa = obtener_o_crear_sesion_pendiente(preinscripcion, usuario_staff_id=current_user.id)
    enlace = enlace_firmables(sesion_activa.token) if (sesion_activa and sesion_activa.token) else None
    return render_template(
        "autoclub/seleccion_femenina/firmables_lista.html",
        preinscripcion=preinscripcion, estado_docs=estado_docs, sesion_activa=sesion_activa, enlace=enlace,
    )


@autoclub_bp.route("/seleccion-femenina/<int:id>/firmables/preparar-enlace", methods=["POST"])
@login_required
def seleccion_femenina_firmables_preparar_enlace(id):
    from app.services.firmables_service import obtener_o_crear_sesion_pendiente
    preinscripcion = PreinscripcionCarcross.query.get_or_404(id)
    obtener_o_crear_sesion_pendiente(preinscripcion, usuario_staff_id=current_user.id)
    flash("Enlace de autorizaciones listo — puedes enviarlo por email, copiarlo o abrirlo aquí mismo.", "success")
    return redirect(url_for("autoclub.seleccion_femenina_firmables", id=id))


@autoclub_bp.route("/seleccion-femenina/<int:id>/firmables/enviar-email", methods=["POST"])
@login_required
def seleccion_femenina_firmables_enviar_email(id):
    from app.services.firmables_service import obtener_o_crear_sesion_pendiente, enviar_email_firmables
    preinscripcion = PreinscripcionCarcross.query.get_or_404(id)
    sesion = obtener_o_crear_sesion_pendiente(preinscripcion, usuario_staff_id=current_user.id)
    if enviar_email_firmables(preinscripcion, sesion):
        registrar_log("cambiar_estado", "preinscripcion_carcross", id,
                      f"Enlace de firmables enviado por email: {preinscripcion.nombre_completo}")
        flash("Enlace enviado por email.", "success")
    else:
        flash("No se pudo enviar el email — revisa el historial de mails o copia el enlace para enviarlo a mano.", "danger")
    return redirect(url_for("autoclub.seleccion_femenina_firmables", id=id))


@autoclub_bp.route("/seleccion-femenina/<int:id>/firmables/historial/<tipo>")
@login_required
def seleccion_femenina_firmables_historial(id, tipo):
    preinscripcion = PreinscripcionCarcross.query.get_or_404(id)
    if tipo not in TIPOS_FIRMABLE_KEYS:
        abort(404)
    versiones = (DocumentoFirmado.query
                 .filter_by(preinscripcion_id=id, tipo=tipo)
                 .order_by(DocumentoFirmado.version.desc()).all())
    return render_template(
        "autoclub/seleccion_femenina/firmables_historial.html",
        preinscripcion=preinscripcion, tipo=tipo, label=TIPOS_FIRMABLE_LABELS[tipo], versiones=versiones,
    )


@autoclub_bp.route("/seleccion-femenina/<int:id>/firmables/documento/<int:doc_id>/pdf")
@login_required
def seleccion_femenina_firmables_pdf(id, doc_id):
    from app.services.firmables_pdf import generar_pdf_documento
    preinscripcion = PreinscripcionCarcross.query.get_or_404(id)
    documento = DocumentoFirmado.query.filter_by(id=doc_id, preinscripcion_id=id).first_or_404()
    pdf_bytes = generar_pdf_documento(preinscripcion, documento)
    return Response(pdf_bytes, mimetype="application/pdf", headers={
        "Content-Disposition": f"inline; filename={documento.tipo}_v{documento.version}.pdf",
    })


# ══════════════════════════════════════════════════════════════════════════
# ENTREVISTA — notas de entrevista por piloto, visibles solo para el staff.
# ══════════════════════════════════════════════════════════════════════════

@autoclub_bp.route("/seleccion-femenina/<int:id>/entrevista")
@login_required
def seleccion_femenina_entrevista(id):
    preinscripcion = PreinscripcionCarcross.query.get_or_404(id)
    entrevista = preinscripcion.entrevista
    return render_template(
        "autoclub/seleccion_femenina/entrevista.html",
        preinscripcion=preinscripcion, entrevista=entrevista,
        bloque1=BLOQUE_1_FAMILIA, bloque2=BLOQUE_2_CANDIDATA,
    )


@autoclub_bp.route("/seleccion-femenina/<int:id>/entrevista/autoguardar", methods=["POST"])
@login_required
def seleccion_femenina_entrevista_autoguardar(id):
    preinscripcion = PreinscripcionCarcross.query.get_or_404(id)
    data = request.get_json(force=True, silent=True) or {}
    notas_entrantes = data.get("notas", {})
    if not isinstance(notas_entrantes, dict):
        return jsonify({"ok": False, "error": "Formato inválido."}), 400
    notas = {k: str(v)[:5000] for k, v in notas_entrantes.items() if k in CLAVES_PREGUNTAS}

    entrevista = preinscripcion.entrevista
    if not entrevista:
        entrevista = EntrevistaPiloto(preinscripcion_id=id)
        db.session.add(entrevista)
    entrevista.notas_json = json.dumps(notas, ensure_ascii=False)
    if data.get("valoracion_bloque1") is not None:
        entrevista.valoracion_bloque1 = str(data.get("valoracion_bloque1"))[:20]
    if data.get("valoracion_bloque2") is not None:
        entrevista.valoracion_bloque2 = str(data.get("valoracion_bloque2"))[:20]
    if not entrevista.fecha_entrevista:
        entrevista.fecha_entrevista = datetime.utcnow()
    entrevista.realizada_por_id = current_user.id
    db.session.commit()
    return jsonify({"ok": True, "guardado_en": entrevista.actualizado_en.strftime("%H:%M:%S")})

